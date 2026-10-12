"""Excel report of one stored run (PLAN.md §7.1, Phase 5).

Built only from what the run stored (work.run, run_version, run_target,
run_input, check_result), so the same run always gives the same report, and a
replay (runner.replay) gives the same report too: load(conn, run_id,
rows=runner.replay(conn, run_id)). Nothing here decides risk: levels come from
the stored results, rule texts from the run's stored rule set (CLAUDE.md rule 4).

Sheets:
  Kopsavilkums  one row per checked object: input, object, addresses, one column
                per module (status / risk), overall risk, warnings
  Detaļas       one row per finding (item), and one per data gap with its reason
  Avoti         sources with version, date, freshness, coverage and processing
                notes; rule set (version, status, checksum) and thresholds
  Neatrasts     input rows (or columns) not checked, with reason and suggestions
Codes are written as text (leading zeros kept, CLAUDE.md rule 8).
"""

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.checks import rules as rules_mod
from engine.checks import runner, summary
from engine.checks.contract import CheckResult
from engine.checks.describe import describe
from engine.checks.modules import BY_ID
from engine.checks.rules import RuleSet
from engine.links import item_links
from engine.map import geo

FOOTER = "Informatīvs pārskats, balstīts uz atvērtajiem datiem; nav oficiāls reģistra izraksts"

STATUS_LV = {
    "found": "Ir",
    "clear": "Nav",
    "no_data": "Nav datu",
    "not_covered": "Nav pārklājuma",
    "not_applicable": "nav piemērojams",
    "source_unavailable": "Avots nav pieejams",
    "error": "Kļūda",
}
RISK_LV = {
    "high": "augsts",
    "medium": "vidējs",
    "unknown": "nezināms",
    "low": "zems",
    "none": "nav",
}
RISK_FILL = {
    "high": "FFC7CE",
    "medium": "FFD9A0",
    "unknown": "D9D9D9",
    "low": "FFF2CC",
    "none": "C6EFCE",
}
NEUTRAL_FILL = "F2F2F2"
TYPE_LV = {
    "property": "Nekustamais īpašums",
    "parcel": "Zemes vienība",
    "building": "Būve",
    "parcel_part": "Zemes vienības daļa",
    "premise_group": "Telpu grupa",
}
INPUT_STATUS_LV = {
    "not_found": "nav atrasts",
    "ambiguous": "neviennozīmīgs",
    "unresolved": "nav nosakāms",
    "outside_pilot": "ārpus pilota teritorijas",
    "invalid": "nederīgs",
    "empty": "tukšs",
    "kind_mismatch": "veids neatbilst",
    "no_objects": "nav kadastra objektu",
    "not_stored": "nav saglabāts",
}
COLUMN_LV = {
    "vertiba": "vērtība",
    "adrese": "adrese",
    "veids": "veids",
    "pieteicejs_reg_nr": "pieteicēja reģ. nr.",
}
RULE_SET_STATUS_LV = {"agreed": "saskaņoti", "draft": "melnraksts"}
THRESHOLD_LV = {
    "inward_buffer_m": "Zemes vienības ģeometrija samazināta uz iekšu (m)",
    "min_overlap_m2": "Minimālais pārklājums atradnei (m²)",
    "min_overlap_share": "Minimālais pārklājums atradnei (daļa no platības)",
    "bis_nearby_m": "BIS objekts „tuvumā” (m)",
}
ADDRESS_UNRESOLVED = "adrese nav atrasta (sk. Neatrasts)"
FROZEN = "frozen_only"

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="305496")
WRAP = Alignment(wrap_text=True, vertical="top")
TEXT = "@"


# --- data ------------------------------------------------------------------------------------


@dataclass
class ReportData:
    head: dict[str, Any]  # summary.run_header
    targets: list[dict[str, Any]]
    inputs: list[dict[str, Any]]
    rows: list[runner.Row]
    rules: RuleSet
    centres: dict[str, tuple[float, float]] = field(default_factory=dict)  # flood viewer links


def load(conn, run_id: int, rows: list[runner.Row] | None = None) -> ReportData | None:
    """Everything the report shows, from the stored run. `rows`: results to show
    instead of the stored ones (a replay of the same run)."""
    head = summary.run_header(conn, run_id)
    if head is None:
        return None
    rule_set = conn.execute("SELECT rule_set FROM work.run WHERE id = %s", (run_id,)).fetchone()[0]
    cols = (
        "target_no",
        "input_no",
        "input",
        "type",
        "designation",
        "label",
        "given_address",
        "project_no",
        "applicant_reg_no",
        "warnings",
    )
    targets = [
        dict(zip(cols, rec, strict=True))
        for rec in conn.execute(
            f"SELECT {', '.join(cols)} FROM work.run_target WHERE run_id = %s"  # noqa: S608
            " ORDER BY target_no",
            (run_id,),
        )
    ]
    rows = runner.sort_rows(rows if rows is not None else runner.stored(conn, run_id))
    return ReportData(
        head=head,
        targets=targets,
        inputs=runner.stored_inputs(conn, run_id),
        rows=rows,
        rules=rules_mod.from_json(rule_set),
        centres=geo.centres(conn, run_id, [r.unit_code for r in rows]),
    )


# --- cells -----------------------------------------------------------------------------------


def _put(ws, row: int, col: int, value: Any, *, text: bool = False, fmt: str | None = None):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # never a formula (input text is data)
    if text:
        cell.number_format = TEXT
    elif fmt:
        cell.number_format = fmt
    cell.alignment = WRAP
    return cell


def _fill(cell, color: str) -> None:
    cell.fill = PatternFill("solid", fgColor=color)


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, int]], first: bool = False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for c, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill, cell.alignment = HEAD_FONT, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"
    ws.oddFooter.center.text = FOOTER
    ws.evenFooter.center.text = FOOTER
    return ws


def _finish(ws, n_columns: int, last_row: int) -> None:
    """Filter over the table, then the footer line below it."""
    ws.auto_filter.ref = f"A1:{get_column_letter(n_columns)}{max(last_row, 1)}"
    cell = ws.cell(row=last_row + 2, column=1, value=FOOTER)
    cell.font = Font(italic=True)


def _num(value: Any) -> float | None:
    return value if isinstance(value, int | float) and not isinstance(value, bool) else None


# --- Kopsavilkums ---------------------------------------------------------------------------


def _modules(data: ReportData) -> list[str]:
    return sorted(data.head["modules"], key=lambda m: BY_ID[m].no)


def _by_target(data: ReportData) -> dict[int, dict[str, CheckResult]]:
    out: dict[int, dict[str, CheckResult]] = {}
    for row in data.rows:
        if row.scope == "target":
            out.setdefault(row.target_no, {})[row.result.module] = row.result
    return out


def _address_unresolved(target: dict[str, Any]) -> bool:
    return any(w.get("code") == "given_address_unresolved" for w in target["warnings"])


def _cell(module: str, r: CheckResult | None, target: dict[str, Any]) -> tuple[str, str]:
    """(text, fill colour) of one module cell."""
    if r is None:
        return "", NEUTRAL_FILL
    b = summary.bucket(module, r.status, r.reason)
    if b == "not_applicable":
        if module == "address_consistency" and _address_unresolved(target):
            return ADDRESS_UNRESOLVED, NEUTRAL_FILL
        return STATUS_LV[b], NEUTRAL_FILL
    return f"{STATUS_LV[b]} / {RISK_LV[r.risk]}", RISK_FILL[r.risk]


def _object_address(results: dict[str, CheckResult]) -> str | None:
    r = results.get("cadastre_facts")
    if r is None:
        return None
    seen: list[str] = []
    for i in r.items:
        a = i.get("address")
        if a and a not in seen:
            seen.append(a)
    return "; ".join(seen) or None


def _given_address(results: dict[str, CheckResult], target, inputs) -> str | None:
    if _address_unresolved(target):
        typed = inputs.get(target["input_no"], {}).get("address")
        return f"nav atrasta: {typed}" if typed else "nav atrasta"
    code = target["given_address"]
    if code is None:
        return None
    r = results.get("address_consistency")
    given = next((i for i in (r.items if r else []) if "given_address" in i), None)
    text = given.get("given_text") if given else None
    return f"{text} ({code})" if text else code


def _warnings(target: dict[str, Any], results: dict[str, CheckResult]) -> str:
    lines: list[str] = []
    frozen = 0

    def add(msg: str) -> None:
        if msg and msg not in lines:
            lines.append(msg)

    for w in target["warnings"]:
        add(w.get("message", ""))
    for r in results.values():
        for w in r.warnings:
            if w.code == FROZEN:
                frozen += 1
                continue
            msg = w.message
            if w.code == "units_partial":  # engine text lists raw status codes
                for code, lv in STATUS_LV.items():
                    msg = msg.replace(f"{code}: ", f"{lv}: ")
            if w.code == "stale_data":  # about a source, not a unit; old runs: age in Latvian
                msg = re.sub(
                    r"\((\d+ [a-z]+(?: \d+ [a-z]+)*)\)$", lambda m: f"({_age(m.group(1))})", msg
                )
                add(msg)
                continue
            unit = getattr(w, "unit", None)
            add(f"{unit}: {msg}" if unit and unit != target["designation"] else msg)
    if frozen:
        add(f"{frozen} apgrūtinājum(i) tikai 2025-01-31 sarakstā – pārbaudīt manuāli (sk. Detaļas)")
    return "\n".join(lines)


def _summary_sheet(wb: Workbook, data: ReportData) -> None:
    modules = _modules(data)
    columns = [
        ("Nr.", 6),
        ("Rinda", 7),
        ("Projekta Nr.", 14),
        ("Ievadītā vērtība", 24),
        ("Pārbaudītais objekts", 18),
        ("Objekta veids", 16),
        ("Interpretācija", 22),
        ("Adrese (kadastrā)", 30),
        ("Norādītā adrese", 30),
        ("Pieteicēja reģ. nr.", 14),
        *[(f"{BY_ID[m].no}. {BY_ID[m].name_lv}", 18) for m in modules],
        ("Kopējais risks", 12),
        ("Brīdinājumi", 60),
    ]
    ws = _sheet(wb, "Kopsavilkums", columns, first=True)
    results = _by_target(data)
    inputs = {i["input_no"]: i for i in data.inputs}
    row = 1
    for t in data.targets:
        row += 1
        res = results.get(t["target_no"], {})
        inp = inputs.get(t["input_no"], {})
        values = [
            (t["target_no"] + 1, False),
            (inp.get("row_no"), False),
            (t["project_no"], True),
            (t["input"], True),
            (t["designation"], True),
            (TYPE_LV.get(t["type"], t["type"]), False),
            (t["label"], False),
            (_object_address(res), False),
            (_given_address(res, t, inputs), False),
            (t["applicant_reg_no"], True),
        ]
        for c, (v, text) in enumerate(values, 1):
            _put(ws, row, c, v, text=text)
        c = len(values)
        applicable = []
        for m in modules:
            c += 1
            r = res.get(m)
            text, color = _cell(m, r, t)
            _fill(_put(ws, row, c, text), color)
            if r is not None and summary.bucket(m, r.status, r.reason) != "not_applicable":
                applicable.append(r.risk)
        worst = data.rules.worst(applicable)
        cell = _put(ws, row, c + 1, RISK_LV[worst] if worst else None)
        if worst:
            _fill(cell, RISK_FILL[worst])
            cell.font = Font(bold=True)
        _put(ws, row, c + 2, _warnings(t, res))
    _finish(ws, len(columns), row)


# --- Detaļas --------------------------------------------------------------------------------


def _address_line(r: CheckResult) -> str:
    given = next((i for i in r.items if "given_address" in i), {})
    linked = [i.get("address") or i.get("varis_code") for i in r.items if "given_address" not in i]
    text = f"Norādītā adrese {given.get('given_text') or given.get('given_address')}"
    text += " nav piesaistīta objektam"
    return text + (f"; objekta adreses kadastrā: {', '.join(linked)}" if linked else "")


def _technical_line(r: CheckResult, rules: RuleSet, designation: str) -> list[dict[str, Any]]:
    """Technical records (rules.yaml not_findings, rule set v2) as one line: listed,
    never a finding, risk none."""
    items = [i for i in r.items if i.get("technical")]
    nf = rules.not_findings.get(r.module)
    if not items or nf is None:
        return []
    codes = sorted({str(i.get("code")) for i in items})
    fixed = {"risk": "none", "ids": [], "texts": ["Nav atradne (riska noteikumi: not_findings)."]}
    text = f"{nf.label_lv}: {len(items)} ({', '.join(codes)}) – nav atradne"
    return [{"text": text, "unit": designation, "fixed": fixed, "item": items[0]}]


def _detail_lines(r: CheckResult, designation: str, rules: RuleSet) -> list[dict[str, Any]]:
    """Lines of one target-level result: one per item (findings), one per data gap,
    one for technical records. "rules": what to explain the line by
    (RuleSet.explain: item or warning); "item": the item its source link is for."""
    b = summary.bucket(r.module, r.status, r.reason)
    technical = _technical_line(r, rules, designation)
    if b in ("clear", "not_applicable"):
        return technical
    lines = []
    if b != "found":
        lines.append({"text": r.reason or STATUS_LV[b], "unit": designation, "rules": None})
        return lines
    if r.module == "address_consistency":
        given = next((i for i in r.items if "given_address" in i), None)
        return [{"text": _address_line(r), "unit": designation, "rules": None, "item": given}]
    frozen_items = set()
    for item in r.items:
        if item.get("technical"):
            continue
        if item.get("frozen"):
            frozen_items.add((item.get("object_code"), item.get("code")))
        lines.append(
            {
                "text": describe(r.module, item),
                "unit": item.get("unit") or item.get("designation") or designation,
                "overlap_m2": _num(item.get("overlap_m2")),
                "share": _num(item.get("share")),
                "match": item.get("match_label") or item.get("match_type") or r.match_type,
                "rules": {"item": item},
                "item": item,
            }
        )
    for w in r.warnings:
        if w.code != FROZEN:
            continue
        key = (getattr(w, "object_code", None), getattr(w, "encumbrance_code", None))
        if key in frozen_items:
            continue
        lines.append(
            {
                "text": f"{key[1]} {getattr(w, 'description', None) or ''}".strip()
                + (f"; reģistrēts {w.establish_date}" if getattr(w, "establish_date", None) else "")
                + f" – {w.message}",
                "unit": getattr(w, "unit", None) or key[0] or designation,
                "rules": {"warning": FROZEN},
                "item": {"object_code": key[0], "code": key[1]},
            }
        )
    return lines + technical


def _details_sheet(wb: Workbook, data: ReportData) -> None:
    columns = [
        ("Nr.", 6),
        ("Pārbaudītais objekts", 18),
        ("Vienība", 18),
        ("Modulis", 22),
        ("Statuss", 14),
        ("Ieraksts", 60),
        ("Pārklājums, m²", 12),
        ("Pārklājums, %", 11),
        ("Atbilstības veids", 24),
        ("Risks", 10),
        ("Noteikums", 26),
        ("Noteikuma teksts", 48),
        ("Avots", 30),
        ("Datu datums", 12),
        ("data.gov.lv saite", 40),
        ("Avots sistēmā", 34),
    ]
    ws = _sheet(wb, "Detaļas", columns)
    targets = {t["target_no"]: t for t in data.targets}
    # A property's units have their own stored results (own rules and risk):
    # show those, not the aggregate's merged lists.
    has_units = {(x.target_no, x.result.module) for x in data.rows if x.scope == "unit"}
    row = 1
    for res in data.rows:
        if res.scope == "target" and (res.target_no, res.result.module) in has_units:
            continue
        r = res.result
        t = targets[res.target_no]
        for line in _detail_lines(r, r.subject.id, data.rules):
            if "fixed" in line:
                ids, texts, risk = (
                    line["fixed"]["ids"],
                    line["fixed"]["texts"],
                    line["fixed"]["risk"],
                )
            else:
                fired = data.rules.explain(r, **(line.get("rules") or {}))
                ids = [f.id for f in fired] or list(r.rule_ids)
                texts = [f.text_lv for f in fired] or [
                    f"Noklusējuma līmenis statusam „{STATUS_LV.get(r.status, r.status)}”"
                ]
                risk = data.rules.worst([f.level for f in fired]) if fired else r.risk
            link = _source_link(r, line, data.centres)
            src = r.source
            row += 1
            values = [
                (res.target_no + 1, None),
                (t["designation"], TEXT),
                (line["unit"], TEXT),
                (f"{r.module_no}. {BY_ID[r.module].name_lv}", None),
                (STATUS_LV[summary.bucket(r.module, r.status, r.reason)], None),
                (line["text"], None),
                (line.get("overlap_m2"), "#,##0.00"),
                (line.get("share"), "0.0%"),
                (line.get("match"), None),
                (RISK_LV[risk], None),
                (", ".join(ids), None),
                ("\n".join(texts), None),
                (src.name if src else None, None),
                (src.dataset_date if src else None, TEXT),
                (src.datagov_url if src else None, None),
                (_link_text(link), None),
            ]
            for c, (v, fmt) in enumerate(values, 1):
                cell = _put(ws, row, c, v, text=fmt == TEXT, fmt=fmt if fmt != TEXT else None)
                if c == 10:
                    _fill(cell, RISK_FILL[risk])
                url = v if c == 15 else (link.url if c == 16 and link else None)
                if url:
                    cell.hyperlink = url
                    cell.font = Font(color="0563C1", underline="single")
    _finish(ws, len(columns), row)


def _source_link(r: CheckResult, line: dict[str, Any], centres: dict) -> Any:
    """The line's "Skatīt avotā" link (engine/links.py), or None for a data gap."""
    item = line.get("item")
    if item is None:
        return None
    unit = line["unit"]
    found = item_links(r.module, item, r.subject.type, unit, centres.get(unit))
    return found[0] if found else None


def _link_text(link: Any) -> str | None:
    if link is None:
        return None
    return f"{link.system} (meklēt: {link.copy})" if link.copy else link.system


# --- Avoti ----------------------------------------------------------------------------------


AGE_UNITS = {"year": "g.", "mon": "mēn.", "day": "d."}


def _age(text: str | None) -> str | None:
    """PostgreSQL interval text ("1 year 1 mon", "14 days") in Latvian."""
    if not text:
        return text
    parts = text.split()
    out = []
    for n, unit in zip(parts[::2], parts[1::2], strict=False):
        out.append(f"{n} {AGE_UNITS.get(unit.rstrip('s'), unit)}")
    return " ".join(out) or text


FRESHNESS_FILL = {
    "current": RISK_FILL["none"],
    "newer_available": RISK_FILL["high"],
    "source_stale": RISK_FILL["medium"],
    "unknown": RISK_FILL["unknown"],
}


def _freshness(v: dict[str, Any]) -> tuple[str, str | None, str | None, str | None, str]:
    """(label, reason, source updated, stale-after limit, fill) of one pinned source.
    Runs before 2026-10-11 stored a max_age verdict instead; shown as it was."""
    f = v.get("freshness")
    if f is not None:
        limit = _age(f.get("source_stale_after")) or "netiek pārbaudīts"
        fill = FRESHNESS_FILL[f["status"]]
        return f["label_lv"], f["text_lv"], f.get("source_updated"), limit, fill
    if "stale" not in v:
        return "nezināms", None, None, None, RISK_FILL["unknown"]
    stale = v["stale"]
    reason = f"dati {'vecāki' if stale else 'jaunāki'} par pieļaujamo vecumu"
    reason += f" ({_age(v.get('max_age'))})"
    fill = RISK_FILL["high"] if stale else RISK_FILL["none"]
    return ("novecojuši" if stale else "aktuāli"), reason, None, None, fill


def _sources_sheet(wb: Workbook, data: ReportData) -> None:
    columns = [
        ("Avots", 34),
        ("Avota id", 22),
        ("Datu kopas versija", 10),
        ("Datu datums", 12),
        ("Svaigums", 14),
        ("Svaiguma pamatojums", 40),
        ("Avots atjaunots", 12),
        ("Brīdināt, ja avots nav atjaunots ilgāk par", 14),
        ("Pārklājums", 40),
        ("Apstrādes piezīmes", 60),
        ("data.gov.lv saite", 40),
    ]
    ws = _sheet(wb, "Avoti", columns)
    row = 1
    for v in data.head["dataset_versions"]:
        row += 1
        label, reason, updated, limit, color = _freshness(v)
        values = [
            v.get("name"),
            v.get("id"),
            v.get("dataset_version_id"),
            v.get("dataset_date"),
            label,
            reason,
            updated,
            limit,
            v.get("coverage"),
            "\n".join(v.get("processing") or []) or None,
            v.get("datagov_url"),
        ]
        for c, value in enumerate(values, 1):
            cell = _put(ws, row, c, value, text=c in (2, 4, 7))
            if c == 5:
                _fill(cell, color)
            if c == len(values) and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
    _finish(ws, len(columns), row)

    rs = data.head["rule_set"]
    created = data.head["created_at"]
    block: list[tuple[str, Any]] = [
        ("Riska noteikumi (config/rules.yaml)", None),
        ("Versija", rs["version"]),
        ("Statuss", RULE_SET_STATUS_LV.get(rs["status"], rs["status"])),
        ("Kontrolsumma (SHA-256)", rs["sha256"]),
        ("Riska secība (sliktākais pirmais)", ", ".join(RISK_LV[r] for r in rs["risk_order"])),
        ("", None),
        ("Sliekšņi", None),
        *[(THRESHOLD_LV.get(k, k), v) for k, v in (rs["thresholds"] or {}).items()],
        ("", None),
        ("Pārbaude", None),
        ("Pārbaudes nr.", data.head["id"]),
        (
            "Izveidota",
            created.strftime("%Y-%m-%d %H:%M %Z") if isinstance(created, datetime) else created,
        ),
        ("Izcelsme", data.head["origin"]),
        ("Moduļi", ", ".join(f"{BY_ID[m].no}. {BY_ID[m].name_lv}" for m in _modules(data))),
        ("Dzinēja versija", data.head["engine_version"]),
        ("", None),
        ("Riska krāsas", None),
    ]
    r = row + 4
    for label, value in block:
        a = _put(ws, r, 1, label)
        if value is None and label:
            a.font = Font(bold=True)
        if value is not None:
            _put(ws, r, 2, value, text=isinstance(value, str))
        r += 1
    for risk in rs["risk_order"]:
        _fill(_put(ws, r, 1, RISK_LV[risk]), RISK_FILL[risk])
        r += 1
    _fill(_put(ws, r, 1, "nav piemērojams / nav atrasta adrese"), NEUTRAL_FILL)


# --- Neatrasts ------------------------------------------------------------------------------


def _suggestions(items: list[dict[str, Any]]) -> str | None:
    lines = []
    for s in items:
        text = f"{s.get('label') or ''} {s.get('designation') or ''}".strip()
        if s.get("description"):
            text += f" – {s['description']}"
        lines.append(text)
    return "\n".join(lines) or None


def _not_found_sheet(wb: Workbook, data: ReportData) -> None:
    columns = [
        ("Rinda", 7),
        ("Kolonna", 16),
        ("Ievadīts", 30),
        ("Projekta Nr.", 14),
        ("Statuss", 18),
        ("Iemesls", 50),
        ("Ieteikumi", 50),
    ]
    ws = _sheet(wb, "Neatrasts", columns)
    row = 1
    for inp in data.inputs:
        for p in inp["problems"]:
            row += 1
            typed = {"vertiba": inp["value"], "adrese": inp["address"]}.get(p["column"])
            values = [
                (inp["row_no"], False),
                (COLUMN_LV.get(p["column"], p["column"]), False),
                (typed, True),
                (inp["project_no"], True),
                (INPUT_STATUS_LV.get(p["status"], p["status"]), False),
                (p.get("reason"), False),
                (_suggestions(p.get("suggestions") or []), False),
            ]
            for c, (v, text) in enumerate(values, 1):
                _put(ws, row, c, v, text=text)
    if row == 1:
        _put(ws, 2, 1, "Visas ievadītās rindas atrastas un pārbaudītas.")
        row = 2
    _finish(ws, len(columns), row)


# --- workbook -------------------------------------------------------------------------------


def build(data: ReportData) -> Workbook:
    wb = Workbook()
    _summary_sheet(wb, data)
    _details_sheet(wb, data)
    _sources_sheet(wb, data)
    _not_found_sheet(wb, data)
    created = data.head["created_at"]
    if isinstance(created, datetime):
        wb.properties.created = wb.properties.modified = created.replace(tzinfo=None)
    wb.properties.title = f"Kadastra pārbaude {data.head['id']}"
    wb.properties.creator = "Kadastra pārbaude"
    return wb


def to_bytes(data: ReportData) -> bytes:
    buf = io.BytesIO()
    build(data).save(buf)
    return buf.getvalue()


def filename(run_id: int) -> str:
    return f"kadastra_parbaude_{run_id}.xlsx"
