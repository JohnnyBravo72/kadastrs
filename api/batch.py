"""Batch runs and the Excel report (PLAN.md §9 Phase 5).

  POST /api/batch                  file (CSV / XLSX) or text (pasted list) -> new run;
                                   format=json (summary + report_url) or xlsx (the report)
  GET  /api/runs/{id}/report.xlsx  the report of a stored run (built from stored results)
  GET  /api/batch/template.xlsx    an empty input template
  GET  /                           a plain page with one form, for colleagues without the CLI

The uploaded file is read in memory and never written by this code (PLAN.md §1);
Starlette may spool an upload over 1 MB to a temporary file that it deletes when
the request ends. Only the cleaned rows and the results are stored.
"""

import io
from html import escape
from pathlib import Path
from typing import Annotated, Any, Literal

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font

from engine.batch import reader
from engine.batch.plan import run_batch
from engine.checks import runner, summary
from engine.resolve.context import SourceUnavailable, resolve_config
from reports import excel

router = APIRouter(tags=["batch"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD = 5 * 1024 * 1024  # 1,000 rows are far below this
PAGE = Path(__file__).resolve().parent / "static" / "index.html"


def _xlsx_response(data: bytes, name: str) -> Response:
    return Response(
        content=data,
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


def _error_page(status: int, message: str) -> HTMLResponse:
    body = (
        "<!doctype html><html lang='lv'><meta charset='utf-8'><title>Kļūda</title>"
        "<body style='font-family:system-ui,sans-serif;max-width:40rem;margin:2rem auto;"
        "padding:0 1rem'><h1>Pārbaudi nevarēja sākt</h1>"
        f"<p>{escape(message)}</p><p><a href='/'>Atpakaļ</a></p></body></html>"
    )
    return HTMLResponse(body, status_code=status)


def _rows(file: UploadFile | None, text: str | None) -> list[reader.InputRow]:
    limit = resolve_config().max_rows
    has_file = file is not None and bool(file.filename)
    if has_file:
        data = file.file.read(MAX_UPLOAD + 1)
        if len(data) > MAX_UPLOAD:
            raise HTTPException(status_code=413, detail="fails ir pārāk liels (vairāk par 5 MB)")
        if data:
            return reader.read_upload(file.filename, data, limit)
    if text and text.strip():
        return reader.read_text(text, limit)
    raise reader.InputError("nav ievadīts saraksts: pievienojiet failu vai ielīmējiet vērtības")


@router.post("/api/batch", status_code=201, response_model=None)
def create_batch(
    request: Request,
    file: Annotated[UploadFile | None, File()] = None,
    text: Annotated[str | None, Form()] = None,
    sources: Annotated[str | None, Form()] = None,
    format: Annotated[Literal["json", "xlsx"], Form()] = "json",  # noqa: A002
) -> Any:
    """A pasted list (`text`) or a CSV / XLSX file (`file`) with the columns
    vertiba, projekta_nr, pieteicejs_reg_nr, adrese, veids; at most 1,000 rows.
    Every row is resolved like POST /api/resolve; rows that are not checked are
    listed on the report's "Neatrasts" sheet."""
    selected = [s.strip() for s in sources.split(",") if s.strip()] if sources else None
    try:
        rows = _rows(file, text)
        with request.app.state.pool.connection() as conn:
            run_id = run_batch(conn, rows, selected)
            if format == "xlsx":
                data = excel.to_bytes(excel.load(conn, run_id))
                return _xlsx_response(data, excel.filename(run_id))
            inputs = runner.stored_inputs(conn, run_id)
            return summary.summarize(conn, run_id) | {
                "inputs": {
                    "rows": len(inputs),
                    "checked": sum(1 for i in inputs if i["target_nos"]),
                    "not_checked": sum(1 for i in inputs if not i["target_nos"]),
                },
                "report_url": f"/api/runs/{run_id}/report.xlsx",
            }
    except (reader.InputError, runner.UnknownSources) as exc:
        status, message = 422, str(exc)
    except SourceUnavailable as exc:
        status, message = 503, str(exc)
    except HTTPException as exc:
        status, message = exc.status_code, str(exc.detail)
    if format == "xlsx":
        return _error_page(status, message)
    raise HTTPException(status_code=status, detail=message)


@router.get("/api/runs/{run_id}/report.xlsx")
def run_report(request: Request, run_id: int) -> Response:
    with request.app.state.pool.connection() as conn:
        data = excel.load(conn, run_id)
        if data is None:
            raise HTTPException(status_code=404, detail=f"pārbaude {run_id} nav atrasta")
        return _xlsx_response(excel.to_bytes(data), excel.filename(run_id))


@router.get("/api/batch/template.xlsx")
def input_template() -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "Saraksts"
    columns = [
        ("vertiba", 30),
        ("projekta_nr", 16),
        ("pieteicejs_reg_nr", 18),
        ("adrese", 40),
        ("veids", 16),
    ]
    for c, (name, width) in enumerate(columns, 1):
        ws.cell(row=1, column=c, value=name).font = Font(bold=True)
        ws.column_dimensions[chr(64 + c)].width = width
        for r in range(2, 1002):  # codes are text: keep leading zeros
            ws.cell(row=r, column=c).number_format = "@"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return _xlsx_response(buf.getvalue(), "saraksta_veidne.xlsx")


@router.get("/", response_class=HTMLResponse, include_in_schema=False)
def page() -> HTMLResponse:
    return HTMLResponse(PAGE.read_text(encoding="utf-8"))
