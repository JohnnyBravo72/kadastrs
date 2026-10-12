"""Compare the engine with the colleagues' manual answers (PLAN.md §9, Phase 4).

Input: a filled tests/data/test_set_template.xlsx (sheet "Testa_kopa"; built by
tests/data/build_templates.py). Each row is one object with the manual answer
per module m1-m6, starting with «Ir: …», «Nav», «Nav datu» or «Nav pārklājuma»,
and the date of the manual check.

The values are resolved like POST /api/resolve, checked in one new run
(origin "compare"), and each object and module is reported as:
  agrees         the manual answer and the engine status are the same
  differs        they are not; the data dates and the check date are shown, so
                 a difference explained by newer data can be told apart
  not_answered   no manual answer (empty, or the template hint left in)
  not_comparable the engine has no comparable status (not resolved, not
                 applicable, source unavailable, error)
Nothing is decided here about risk; the engine's risk is only shown.
"""

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine.batch.headers import canonical, fold
from engine.batch.plan import ACCEPTS, KINDS
from engine.checks import runner, summary
from engine.checks.contract import TargetIn, Warning_
from engine.resolve import service
from engine.resolve.context import live_versions, pilot, resolve_config
from engine.resolve.places import load_places

SHEET = "Testa_kopa"
MODULE_COLUMNS = {
    "m1_kadastra_dati": "cadastre_facts",
    "m2_adreses_atbilstiba": "address_consistency",
    "m3_aizsargajama_daba": "nature",
    "m4_pludi_10_gadi": "flood_10y",
    "m5_buvniecibas_lietas": "bis",
    "m6_apgrutinajumi": "encumbrances",
}
COMPARABLE = {"found", "clear", "no_data", "not_covered"}
STATUS_LV = {
    "found": "Ir",
    "clear": "Nav",
    "no_data": "Nav datu",
    "not_covered": "Nav pārklājuma",
    "not_applicable": "Neattiecas",
    "source_unavailable": "Avots nav pieejams",
    "error": "Kļūda",
}
OUTCOME_LV = {
    "agrees": "sakrīt",
    "differs": "atšķiras",
    "not_answered": "nav atbildes",
    "not_comparable": "nav salīdzināms",
}


def _fold(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c)).lower().strip()


def parse_answer(value: Any) -> str | None:
    """Manual answer -> engine status, or None when not answered."""
    if value is None:
        return None
    text = _fold(str(value))
    if not text or text.startswith("‹") or text.startswith("<"):
        return None
    words = re.findall(r"[a-z]+", text)
    if words[:2] == ["nav", "parklajuma"]:
        return "not_covered"
    if words[:2] == ["nav", "datu"]:
        return "no_data"
    if words[:1] == ["nav"]:
        return "clear"
    if words[:1] == ["ir"]:
        return "found"
    return None


def _text(value: Any) -> str | None:
    """Cell value as text; codes are never numbers (CLAUDE.md rule 8)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() or None


@dataclass
class ManualRow:
    row_no: int  # Excel row
    project_no: str | None
    value: str | None
    applicant: str | None
    answers: dict[str, Any]  # module id -> raw answer
    checked_on: date | None
    notes: str | None
    kind: str | None = None  # veids: property | parcel | building (removes ambiguity)


def read_template(path: Path) -> list[ManualRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    # new Latvian headers or the old technical names (engine/batch/headers.py)
    col = {}
    for i, h in enumerate(next(rows)):
        name = canonical(h)
        if name and name not in col:
            col[name] = i
    missing = [c for c in ["projekta_nr", "vertiba", *MODULE_COLUMNS] if c not in col]
    if missing:
        raise ValueError(f"{path.name}: sheet {SHEET} lacks columns {missing}")
    out = []
    for n, cells in enumerate(rows, start=2):

        def get(name, cells=cells):
            i = col.get(name)
            return cells[i] if i is not None and i < len(cells) else None

        value = _text(get("vertiba"))
        project = _text(get("projekta_nr"))
        if value is None or (project or "").upper().startswith("PIEMĒRS"):
            continue
        checked = get("parbaudes_datums")
        if isinstance(checked, datetime):
            checked = checked.date()
        elif isinstance(checked, str):
            try:
                checked = date.fromisoformat(checked.strip())
            except ValueError:
                checked = None
        out.append(
            ManualRow(
                row_no=n,
                project_no=project,
                value=value,
                applicant=_text(get("pieteicejs_reg_nr")),
                answers={m: get(c) for c, m in MODULE_COLUMNS.items()},
                checked_on=checked if isinstance(checked, date) else None,
                notes=_text(get("avots_piezimes")),
                kind=KINDS.get(fold(_text(get("veids")) or "")),
            )
        )
    return out


@dataclass
class Resolved:
    row: ManualRow
    target: TargetIn | None
    note: str | None = None  # why not resolved, or which object was taken


def resolve_rows(conn, rows: list[ManualRow]) -> list[Resolved]:
    v = live_versions(conn)
    p = pilot()
    res = service.resolve(
        conn,
        v,
        load_places(conn, v, p),
        p,
        resolve_config(),
        [service.RowIn(value=r.value) for r in rows],
    )
    out = []
    for row, rr in zip(rows, res.rows, strict=True):
        if rr.status not in ("resolved", "ambiguous") or not rr.matches:
            out.append(Resolved(row, None, f"{rr.status}: {rr.reason or ''}".strip(": ")))
            continue
        m = rr.matches[0]
        if row.kind:
            m = next((x for x in rr.matches if x.type in ACCEPTS[row.kind]), m)
        notes = [Warning_(code=w.code, message=w.message) for w in rr.warnings]
        notes += [Warning_(code=c.code, message=c.message) for c in rr.corrections]
        note = None
        if rr.status == "ambiguous" and not (row.kind and m.type in ACCEPTS[row.kind]):
            note = f"neviennozīmīgs; pārbaudīts kā {m.label} {m.designation}"
        if m.type == "address":
            objects = sorted(m.objects, key=lambda o: (o.type != "parcel", o.designation))
            if not objects:
                out.append(Resolved(row, None, "adresei nav kadastra objektu"))
                continue
            obj = objects[0]
            if len(objects) > 1:
                note = f"adresei {len(objects)} objekti; pārbaudīts {obj.designation}"
            code = m.address.varis_code if m.address else m.designation
            target = TargetIn(
                type=obj.type,
                designation=obj.designation,
                input=row.value,
                given_address=code,
                project_no=row.project_no,
                warnings=notes,
            )
        else:
            target = TargetIn(
                type=m.type,
                designation=m.designation,
                input=row.value,
                project_no=row.project_no,
                warnings=notes,
            )
        out.append(Resolved(row, target, note))
    return out


@dataclass
class Line:
    row_no: int
    project_no: str | None
    value: str | None
    designation: str | None
    module: str
    manual_raw: str | None
    manual: str | None
    engine: str | None  # summary status (not_applicable included)
    outcome: str
    risk: str | None = None
    detail: str | None = None
    data_dates: str | None = None
    checked_on: date | None = None
    newer_data: bool | None = None  # a source date after the manual check
    note: str | None = None


@dataclass
class Report:
    run_id: int | None
    rule_set: str
    lines: list[Line] = field(default_factory=list)

    def counts(self) -> dict[str, dict[str, int]]:
        out: dict[str, dict[str, int]] = {}
        for line in self.lines:
            m = out.setdefault(line.module, dict.fromkeys(OUTCOME_LV, 0))
            m[line.outcome] += 1
        return out


def _detail(module: str, r) -> str | None:
    if r is None:
        return None
    if module == "nature":
        return (
            "; ".join(
                f"{i['kind_lv']} {i.get('name') or ''} ({i['share']:.1%})".strip() for i in r.items
            )
            or None
        )
    if module == "flood_10y":
        return "; ".join(f"{i['overlap_m2']} m² ({i['share']:.1%})" for i in r.items) or r.reason
    if module == "bis":
        return "; ".join(f"{i['case_number']} ({i['match_type']})" for i in r.items) or None
    if module == "encumbrances":
        codes = [i["code"] for i in r.items if not i.get("frozen")]
        frozen = [w.encumbrance_code for w in r.warnings if w.code == "frozen_only"]
        parts = [", ".join(codes)] if codes else []
        if frozen:
            parts.append("tikai 2025-01-31 sarakstā: " + ", ".join(frozen))
        return "; ".join(parts) or None
    if module == "address_consistency":
        return r.reason or r.match_type
    if module == "cadastre_facts" and r.items:
        i = r.items[0]
        return f"{i.get('type')} {i.get('area_m2') or ''} m², {i.get('address') or ''}".strip()
    return r.reason


def compare(conn, path: Path, sources: list[str] | None = None) -> Report:
    rows = read_template(path)
    resolved = resolve_rows(conn, rows)
    targets = [r.target for r in resolved if r.target is not None]
    run_id = runner.start_run(conn, targets, sources, origin="compare") if targets else None
    head = summary.run_header(conn, run_id) if run_id else None
    rule_set = f"{head['rule_set']['version']} ({head['rule_set']['status']})" if head else ""
    results: dict[tuple[int, str], Any] = {}
    if run_id:
        for row in runner.stored(conn, run_id):
            if row.scope == "target":
                results[(row.target_no, row.result.module)] = row.result
    report = Report(run_id, rule_set)
    target_no = 0
    for res in resolved:
        no = None
        if res.target is not None:
            no, target_no = target_no, target_no + 1
        for module in MODULE_COLUMNS.values():
            raw = res.row.answers.get(module)
            manual = parse_answer(raw)
            r = results.get((no, module)) if no is not None else None
            engine = summary.bucket(module, r.status, r.reason) if r else None
            if manual is None:
                outcome = "not_answered"
            elif engine not in COMPARABLE:
                outcome = "not_comparable"
            else:
                outcome = "agrees" if engine == manual else "differs"
            dates = sorted({s.dataset_date for s in r.sources}) if r else []
            newer = (
                any(date.fromisoformat(d) > res.row.checked_on for d in dates)
                if dates and res.row.checked_on
                else None
            )
            report.lines.append(
                Line(
                    row_no=res.row.row_no,
                    project_no=res.row.project_no,
                    value=res.row.value,
                    designation=res.target.designation if res.target else None,
                    module=module,
                    manual_raw=_text(raw),
                    manual=manual,
                    engine=engine,
                    outcome=outcome,
                    risk=r.risk if r else None,
                    detail=_detail(module, r),
                    data_dates=", ".join(dates) or None,
                    checked_on=res.row.checked_on,
                    newer_data=newer,
                    note=res.note,
                )
            )
    return report


# --- output ---------------------------------------------------------------------------------

COLUMNS = [
    ("Rinda", 7),
    ("Projekts", 14),
    ("Ievadīts", 22),
    ("Pārbaudīts objekts", 20),
    ("Modulis", 22),
    ("Manuālā atbilde", 30),
    ("Manuāli", 14),
    ("Rīks", 16),
    ("Rezultāts", 14),
    ("Risks (rīks)", 11),
    ("Rīka atradne", 50),
    ("Datu datumi", 24),
    ("Pārbaudes datums", 14),
    ("Dati jaunāki par pārbaudi", 12),
    ("Piezīme", 40),
]
FILL = {"agrees": "C6EFCE", "differs": "FFC7CE", "not_comparable": "FFEB9C"}


def write_xlsx(report: Report, out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Salidzinajums"
    for c, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    for r, line in enumerate(report.lines, 2):
        values = [
            line.row_no,
            line.project_no,
            line.value,
            line.designation,
            line.module,
            line.manual_raw,
            STATUS_LV.get(line.manual or "", line.manual),
            STATUS_LV.get(line.engine or "", line.engine),
            OUTCOME_LV[line.outcome],
            line.risk,
            line.detail,
            line.data_dates,
            line.checked_on,
            None if line.newer_data is None else ("jā" if line.newer_data else "nē"),
            line.note,
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in (2, 3, 4):
                cell.number_format = "@"  # codes are text (CLAUDE.md rule 8)
            if c == 13:
                cell.number_format = "yyyy-mm-dd"
        if line.outcome in FILL:
            ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor=FILL[line.outcome])
    s = wb.create_sheet("Kopsavilkums")
    s.append(["Pārbaude (run)", report.run_id])
    s.append(["Riska noteikumi", report.rule_set])
    s.append([])
    s.append(["Modulis", *OUTCOME_LV.values()])
    for module, counts in report.counts().items():
        s.append([module, *counts.values()])
    s.column_dimensions["A"].width = 24
    wb.save(out)


def print_report(report: Report) -> None:
    print(f"run {report.run_id}, rule set {report.rule_set}")
    print(f"{'module':22} " + " ".join(f"{o:>15}" for o in OUTCOME_LV))
    for module, counts in report.counts().items():
        print(f"{module:22} " + " ".join(f"{counts[o]:>15}" for o in OUTCOME_LV))
    for line in report.lines:
        if line.outcome == "differs":
            hint = " (dati jaunāki par pārbaudi)" if line.newer_data else ""
            print(
                f"  rinda {line.row_no} {line.value} {line.module}: manuāli {line.manual}, "
                f"rīks {line.engine}{hint} – {line.detail or ''}"
            )
