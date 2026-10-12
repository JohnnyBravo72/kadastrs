"""Batch input (PLAN.md §8.1 step 6, Phase 5): a pasted list, a CSV or an XLSX file
in, one InputRow per row out. Nothing is written to disk; the caller keeps only
the cleaned rows (CLAUDE.md rule 5, PLAN.md §1 "uploads are not stored").

Columns: "Kadastra Nr., apzīmējums vai adrese" (required), and optionally
"Projekta Nr.", "Pieteicēja reģ. Nr.", "Adrese pieteikumā", "Objekta veids"
(īpašums | zemes vienība | būve); the old names `vertiba`, `projekta_nr`,
`pieteicejs_reg_nr`, `adrese`, `veids` work too (engine/batch/headers.py).
Without a value header the first column is the value; a pasted list without a
header has one value per line.

CSV: UTF-8 (with or without BOM), else Windows-1257; delimiter `;` or `,`.
XLSX: the sheet "Testa_kopa" (test-set template) or "Saraksts" (input template)
if present, else the first sheet. Value cells keep their Excel type (a number
that lost its leading zero is repaired and flagged by the classifier, §5.1);
other code columns are read as
text (CLAUDE.md rule 8). Rows whose projekta_nr starts with "PIEMĒRS" are the
template's examples and are skipped.
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from engine.batch.headers import canonical

COLUMNS = {
    "vertiba": "value",
    "projekta_nr": "project_no",
    "pieteicejs_reg_nr": "applicant",
    "adrese": "address",
    "veids": "kind",
}
TEMPLATE_SHEETS = ("Testa_kopa", "Saraksts")  # test-set template, input template
EXAMPLE_PREFIX = "PIEMĒRS"


class InputError(ValueError):
    """The input cannot be read (message in Latvian, shown to the user)."""


@dataclass
class InputRow:
    row_no: int  # 1-based row in the file or pasted text (header included)
    value: str | int | float | None
    kind: str | None = None  # veids as written
    project_no: str | None = None
    address: str | None = None
    applicant: str | None = None


def text_of(value: Any) -> str | None:
    """A cell as text; codes are never numbers (CLAUDE.md rule 8)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime | date):
        value = value.isoformat()
    return str(value).strip() or None


def _blank(cells: list[Any]) -> bool:
    return all(text_of(c) is None for c in cells)


def _header(cells: list[Any]) -> dict[str, int] | None:
    """Column positions if the row is a header with a value column (new Latvian
    or old technical names)."""
    names = [canonical(c) for c in cells]
    if "vertiba" not in names:
        return None
    return {COLUMNS[n]: i for i, n in reversed(list(enumerate(names))) if n in COLUMNS}


def _rows(table: list[tuple[int, list[Any]]], max_rows: int) -> list[InputRow]:
    """Rows (row_no, cells) -> InputRows: header detection, blanks and examples skipped."""
    table = [(n, cells) for n, cells in table if not _blank(cells)]
    if not table:
        return []
    col = _header(table[0][1])
    if col is not None:
        table = table[1:]
    else:
        col = {"value": 0}
    out = []
    for n, cells in table:

        def get(name, cells=cells):
            i = col.get(name)
            return cells[i] if i is not None and i < len(cells) else None

        project = text_of(get("project_no"))
        if project and project.upper().startswith(EXAMPLE_PREFIX):
            continue
        value = get("value")
        if isinstance(value, str):
            value = value.strip() or None
        elif not isinstance(value, int | float) or isinstance(value, bool):
            value = text_of(value)
        out.append(
            InputRow(
                row_no=n,
                value=value,
                kind=text_of(get("kind")),
                project_no=project,
                address=text_of(get("address")),
                applicant=text_of(get("applicant")),
            )
        )
    if len(out) > max_rows:
        raise InputError(f"ne vairāk kā {max_rows} rindas vienā pārbaudē (ir {len(out)})")
    return out


# --- text and CSV ---------------------------------------------------------------------------


def decode(data: bytes) -> str:
    """UTF-8 (BOM or not) wins; otherwise Windows-1257 (Baltic)."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1257")


def _delimiter(lines: list[str]) -> str | None:
    """Delimiter of a table with a `vertiba` header, or of a headerless file
    (`;` or tab only: an address contains commas)."""
    first = next((ln for ln in lines if ln.strip()), "")
    # a header row: the delimiter that splits it into the most known columns
    # (the new value header itself contains a comma)
    known = {}
    for d in ("\t", ";", ","):
        names = [canonical(c) for c in next(csv.reader([first], delimiter=d))]
        if "vertiba" in names:
            known[d] = sum(1 for n in names if n)
    if known:
        best = max(known, key=lambda d: known[d])
        return best if any(d in first for d in ("\t", ";", ",")) else None
    for d in ("\t", ";"):
        if d in first:
            return d
    return None


def read_text(text: str, max_rows: int) -> list[InputRow]:
    """Pasted text: one value per line, or a table with a header (tab, `;` or `,`)."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    d = _delimiter(lines)
    if d is None:
        table = [(n, [ln.strip()]) for n, ln in enumerate(lines, 1)]
    else:
        reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=d)
        table = [(n, cells) for n, cells in enumerate(reader, 1)]
    return _rows(table, max_rows)


def read_csv(data: bytes, max_rows: int) -> list[InputRow]:
    return read_text(decode(data), max_rows)


# --- XLSX -----------------------------------------------------------------------------------


def read_xlsx(data: bytes, max_rows: int) -> list[InputRow]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001  (any unreadable file is the same error)
        raise InputError(f"Excel failu nevar nolasīt: {type(exc).__name__}") from exc
    try:
        named = [n for n in TEMPLATE_SHEETS if n in wb.sheetnames]
        ws = wb[named[0]] if named else wb.worksheets[0]
        table = []
        for n, cells in enumerate(ws.iter_rows(values_only=True), 1):
            cells = list(cells)
            if _blank(cells):
                continue
            table.append((n, cells))
            # header and template examples are not counted (exact check in _rows);
            # stop reading a far too long sheet, never cut it silently
            if len(table) > max_rows + 50:
                raise InputError(f"ne vairāk kā {max_rows} rindas vienā pārbaudē")
        return _rows(table, max_rows)
    finally:
        wb.close()


def read_upload(filename: str | None, data: bytes, max_rows: int) -> list[InputRow]:
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise InputError("vecais .xls formāts netiek atbalstīts; saglabājiet kā .xlsx vai .csv")
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return read_xlsx(data, max_rows)
    return read_csv(data, max_rows)
