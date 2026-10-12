"""Excel templates for colleagues (decided 2026-10-12): readable Latvian headers
(engine/batch/headers.py), required columns dark and optional ones light, a
note on each header, code columns as text (leading zeros kept, CLAUDE.md rule 8)
and the object type as a drop-down list.

- input_template(): the list for a check (GET /api/batch/template.xlsx, and the
  link on the input page); sheet "Saraksts" first, "Instrukcija" second.
- style_headers() / instructions(): shared with tests/data/build_templates.py
  (the test-set template).
"""

from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from engine.batch.headers import INPUT_COLUMNS, KIND_CHOICES, Column

REQUIRED_FILL = PatternFill("solid", fgColor="1F3A5F")
OPTIONAL_FILL = PatternFill("solid", fgColor="D9E2EF")
REQUIRED_FONT = Font(bold=True, color="FFFFFF")
OPTIONAL_FONT = Font(bold=True, color="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="center")
ROWS = 1000  # formatted rows (the check accepts at most 1,000)

NOTES = {
    "vertiba": "Obligāta. Kadastra numurs (11 cipari), apzīmējums (11, 14, 15 vai 17 cipari)"
    " vai adrese.",
    "projekta_nr": "Neobligāta. Projekta numurs (brīvs teksts); grupē objektus pārskatā.",
    "pieteicejs_reg_nr": "Neobligāta. Tikai juridiskas personas reģistrācijas numurs"
    " (11 cipari, sākas ar 4, 5 vai 9). Personas kodus neievadiet: tie netiek saglabāti.",
    "adrese": "Neobligāta. Adrese, kā norādīta pieteikumā: rīks pārbauda, vai tā atbilst objektam.",
    "veids": "Neobligāta. Izvēlieties, ja 11 ciparu numurs var būt gan īpašums, gan citam"
    " īpašumam piederoša zemes vienība; bez tā rīks pārbauda abus.",
}
TEXT_COLUMNS = {"vertiba", "projekta_nr", "pieteicejs_reg_nr", "adrese", "veids"}


def style_headers(ws: Worksheet, columns: Sequence[tuple[Column, int]], rows: int = ROWS) -> None:
    """Header row 1 (required dark, optional light, note on each), widths, code
    columns as text, the object type as a list, the header row frozen."""
    for i, (col, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.fill = REQUIRED_FILL if col.required else OPTIONAL_FILL
        cell.font = REQUIRED_FONT if col.required else OPTIONAL_FONT
        cell.alignment = WRAP
        note = NOTES.get(col.name) or ("Obligāta." if col.required else "Neobligāta.")
        cell.comment = Comment(note, "Kadastra pārbaude", width=320, height=110)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if col.name in TEXT_COLUMNS:
            for r in range(2, rows + 2):
                ws[f"{letter}{r}"].number_format = "@"
        if col.name == "veids":
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(KIND_CHOICES) + '"',
                allow_blank=True,
                error="Izvēlieties: " + ", ".join(KIND_CHOICES) + ".",
                errorTitle="Objekta veids",
            )
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{rows + 1}")
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "A2"


def instructions(ws: Worksheet, lines: Sequence[str]) -> None:
    ws.column_dimensions["A"].width = 120
    for r, text in enumerate(lines, 1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1:
            cell.font = Font(bold=True, size=14)
        elif text.endswith(":"):
            cell.font = Font(bold=True)


LEGEND = [
    "Kolonnu krāsas: tumši zila = obligāta, gaiši zila = neobligāta. Virsrakstu piezīmēs"
    " (sarkanais stūrītis) – paskaidrojums.",
    "Kolonnu nosaukumus nemainiet; der arī vecie nosaukumi (vertiba, projekta_nr,"
    " pieteicejs_reg_nr, adrese, veids).",
]

INPUT_INSTRUCTIONS = [
    "Kadastra pārbaude: saraksta veidne",
    "Aizpildiet lapu „Saraksts”: viena rinda = viens objekts. Augšupielādējiet failu"
    " pārbaudes lapā (http://127.0.0.1:8000/) vai ielīmējiet rindas kopā ar virsrakstu.",
    "",
    "Kolonnas:",
    f"• {INPUT_COLUMNS[0].header} (obligāta): kadastra numurs (11 cipari), apzīmējums"
    " (11, 14, 15 vai 17 cipari) vai adrese. Ciparus ievadiet ar sākuma nullēm"
    " (piem., 05000041006); kolonna ir teksta formātā.",
    f"• {INPUT_COLUMNS[1].header}: projekta numurs; objekti pārskatā tiek grupēti pēc tā.",
    f"• {INPUT_COLUMNS[2].header}: tikai juridiskas personas reģistrācijas numurs. Fizisko"
    " personu vārdus un personas kodus neievadiet: tie netiek saglabāti.",
    f"• {INPUT_COLUMNS[3].header}: adrese, kā norādīta pieteikumā; rīks pārbauda, vai tā"
    " atbilst objektam (modulis „Adreses atbilstība”).",
    f"• {INPUT_COLUMNS[4].header}: izvēlieties no saraksta, ja 11 ciparu numurs ir gan"
    " īpašuma numurs, gan citam īpašumam piederošas zemes vienības apzīmējums. Bez tā rīks"
    " pārbauda abus.",
    "",
    *LEGEND,
    "Ne vairāk kā 1000 rindu vienā pārbaudē. Augšupielādētais fails netiek saglabāts.",
]


def input_template() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Saraksts"
    widths = (36, 16, 26, 36, 30)
    style_headers(ws, list(zip(INPUT_COLUMNS, widths, strict=True)))
    instructions(wb.create_sheet("Instrukcija"), INPUT_INSTRUCTIONS)
    return wb
