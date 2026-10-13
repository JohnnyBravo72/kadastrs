"""Phase 8b importers (PLAN.md §6.1 modules 10–13, 18; docs/decisions.md D11).

  vid_rating           VID taxpayer rating (name never stored)
  deminimis            FM live CSV export via D1 E1 (legal entities only)
  eu_projects          three CFLA periods: projects, locations, partners
  eis_procurement      EIS yearly result files (winner name never stored)
  fid_frozen_property  FID TA, VZD frozen real estate (no person columns)

Personal-data rules (CLAUDE.md rule 5, D11): legal-entity reg. nos only where
applicable; natural-person / IK EU rows keep only project and cadastre fields;
TLS verification stays on (deminimis uses the pinned Sectigo intermediate).
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from etl.identity import FileRecord, iso
from etl.importers.common import copy_rows, fetch_whole, resource_date
from etl.importers.ur import legal_reg_no
from etl.remote import download
from etl.schema_check import SchemaMismatch, check_columns
from etl.settings import schemas
from etl.textio import clean_text, csv_header, iter_csv
from etl.tls import ssl_context_with_sectigo

_ISO = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}")
_DMY = re.compile(r"^([0-9]{2})\.([0-9]{2})\.([0-9]{4})$")
_REGUL_KEY = re.compile(r"(20[0-9]{2}/[0-9]+|[0-9]+/[0-9]{4})")
NATURAL_TYPES = {
    "Fiziska persona",
    "Fiziska persona, kas ir reģistrējusies kā saimnieciskās darbības veicēja",
    "Individuālais komersants",
}


def _iso_date(value: str | None, what: str) -> date | None:
    value = (value or "").strip()
    if not value:
        return None
    if _ISO.match(value):
        return date.fromisoformat(value[:10])
    m = _DMY.match(value)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    raise ValueError(f"{what}: date {value!r} is not YYYY-MM-DD or DD.MM.YYYY")


def _decimal(value: str | None, what: str) -> Decimal | None:
    value = (value or "").strip().replace(" ", "").replace(",", ".")
    if not value:
        return None
    try:
        return Decimal(value)
    except InvalidOperation as exc:
        raise ValueError(f"{what}: not a number: {value!r}") from exc


def _cadastre_text(value) -> str | None:
    """Cadastre number as text; restore a leading zero lost when Excel stored a number (E2)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        text = str(value)
        for length in (11, 14, 15, 17):
            if len(text) == length - 1:
                return "0" + text
        return text
    text = re.sub(r"\s+", "", str(value).strip())
    return text or None


def _regul_key(title: str) -> str:
    text = title or ""
    low = text.casefold()
    # road-freight sub-ceiling under 1407/2013 (approved Phase 8b plan)
    if "1407/2013" in text and (
        "road" in low or "freight" in low or "kravu" in low or "autopārvad" in low
    ):
        return "1407/2013-road"
    m = _REGUL_KEY.search(text)
    return m.group(1) if m else text.strip() or "unknown"


# --- VID ----------------------------------------------------------------------

VID_COLUMNS = ["dataset_version_id", "reg_no", "rating", "explanation", "rated_on"]


class VidRating:
    source_id = "vid_rating"
    FILE = "vid_rating.csv"

    def check_schema(self, ctx):
        header = csv_header(ctx.workdir / self.FILE, delimiter=",")
        check_columns(self.FILE, header, schemas()["csv"]["vid_rating"])

    def fetch(self, ctx):
        files = [fetch_whole(ctx, "rating", self.FILE)]
        self.check_schema(ctx)
        return files

    def version_date(self, ctx, files):
        return resource_date(ctx, ["rating"])

    def load(self, ctx, conn, version_id, files):
        self.check_schema(ctx)
        counts: Counter = Counter()

        def rows():
            for line, r in enumerate(iter_csv(ctx.workdir / self.FILE, delimiter=","), start=2):
                counts["source_rows"] += 1
                reg = legal_reg_no(r["Registracijas_kods"])
                if reg is None:
                    counts["skipped_reg_no_shape"] += 1
                    continue
                rating = (r["Reitings"] or "").strip()
                if rating not in {"A", "B", "C", "J", "N"}:
                    raise ValueError(f"vid line {line}: rating {rating!r}")
                # Nosaukums is never read into a stored column (D11)
                counts["loaded"] += 1
                yield (
                    version_id,
                    reg,
                    rating,
                    clean_text(r["Skaidrojums"]),
                    _iso_date(r["Informacijas_atjaunosanas_datums"], f"line {line}"),
                )

        loaded = copy_rows(conn, "ref.vid_rating", VID_COLUMNS, rows())
        ctx.processing = [
            {
                "id": "personal_data",
                "what": "Name column never stored; non-legal-entity reg. nos skipped (D11)",
                "skipped_reg_no_shape": counts["skipped_reg_no_shape"],
            }
        ]
        return {"vid_rating": {"source_rows": counts["source_rows"], "loaded": loaded}}


# --- De minimis ---------------------------------------------------------------

DEMINIMIS_COLUMNS = [
    "dataset_version_id",
    "row_no",
    "reg_no",
    "assign_date",
    "assigner",
    "programme",
    "amount",
    "regulation",
    "regulation_key",
    "instrument",
]


class Deminimis:
    source_id = "deminimis"
    FILE = "deminimis.csv"

    def _e1_url(self, ctx) -> str:
        e1 = next((e for e in ctx.registry.allowed_external_resources if e.id == "E1"), None)
        if e1 is None or e1.status != "allowed":
            raise LookupError("D1 E1 is not an allowed external resource")
        return e1.urls[0]

    def check_schema(self, ctx):
        header = csv_header(ctx.workdir / self.FILE, delimiter=",")
        check_columns(self.FILE, header, schemas()["csv"]["deminimis"])

    def fetch(self, ctx):
        import httpx

        url = self._e1_url(ctx)
        # dedicated client with the pinned intermediate (CLAUDE.md rule 9, D11)
        with httpx.Client(
            headers=ctx.client.headers,
            follow_redirects=True,
            timeout=ctx.client.timeout,
            verify=ssl_context_with_sectigo(),
        ) as client:
            info = download(client, url, ctx.workdir / self.FILE)
        ctx.source_last_modified = datetime.now(UTC).replace(tzinfo=None)
        self.check_schema(ctx)
        return [
            FileRecord(
                role="export",
                resource_id="E1",
                url=url,
                last_modified=iso(ctx.source_last_modified),
                size=info.size,
                sha256=info.sha256,
                path=self.FILE,
            )
        ]

    def version_date(self, ctx, files):
        # file's own date (E7): newest assign_date (D11)
        latest: date | None = None
        for _, r in iter_csv(ctx.workdir / self.FILE, delimiter=","):
            d = _iso_date(r["assign_date"], "assign_date")
            if d and (latest is None or d > latest):
                latest = d
        if latest is None:
            raise ValueError("deminimis: no assign_date in file")
        return latest

    def load(self, ctx, conn, version_id, files):
        self.check_schema(ctx)
        counts: Counter = Counter()

        def rows():
            for line, r in enumerate(iter_csv(ctx.workdir / self.FILE, delimiter=","), start=2):
                counts["source_rows"] += 1
                reg = legal_reg_no(r["person_reg_nr"])
                if reg is None:
                    counts["skipped_not_legal_entity"] += 1
                    continue
                # person_title and project_title never read into stored columns
                amount = _decimal(r["assigned_amount"], f"line {line}")
                if amount is None:
                    raise ValueError(f"deminimis line {line}: empty amount")
                regul = clean_text(r["regul_title"]) or ""
                counts["loaded"] += 1
                yield (
                    version_id,
                    counts["loaded"],
                    reg,
                    _iso_date(r["assign_date"], f"line {line}"),
                    clean_text(r["assigner_title"]),
                    clean_text(r["program_title"]),
                    amount,
                    regul,
                    _regul_key(regul),
                    clean_text(r["instrument_title"]),
                )

        loaded = copy_rows(conn, "ref.deminimis_aid", DEMINIMIS_COLUMNS, rows())
        ctx.processing = [
            {
                "id": "personal_data",
                "what": "Legal-entity rows only; person_title and project_title never stored (D11)",
                "skipped_not_legal_entity": counts["skipped_not_legal_entity"],
            }
        ]
        return {"deminimis_aid": {"source_rows": counts["source_rows"], "loaded": loaded}}


# --- EU projects --------------------------------------------------------------

# key -> (period, schema, start_col, end_col, fund_col, total_col, eu_col)
PERIODS = {
    "2127": (
        "2021-27",
        "eu_projects_2127",
        "ProjektaIeviesanasSakumaDatums",
        "ProjektaIeviesanasBeiguDatums",
        "Fonds",
        "ProjektaKopejasIzmaksas",
        "ProjektamPieskirtaisESFonduLidzfinansejums",
    ),
    "1420": (
        "2014-20",
        "eu_projects_1420",
        "SakumaDatums",
        "BeiguDatums",
        "EsFonds",
        "KopejaisFinansejums",
        "EsFonduLidzfinansejums",
    ),
    "rrf": (
        "RRF",
        "eu_projects_rrf",
        "SakumaDatums",
        "BeiguDatums",
        "EsFonds",
        "KopejaisFinansejums",
        "EsFonduLidzfinansejums",
    ),
}

EU_PROJECT_COLUMNS = [
    "dataset_version_id",
    "period",
    "project_no",
    "title",
    "status",
    "applicant_type",
    "applicant_reg_no",
    "applicant_name",
    "is_natural_or_ik",
    "started_on",
    "ended_on",
    "fund",
    "measure_no",
    "measure_name",
    "total_cost",
    "eu_funding",
    "monitoring_years",
    "updated_on",
]
EU_LOC_COLUMNS = [
    "dataset_version_id",
    "period",
    "project_no",
    "row_no",
    "cadastre_nr",
    "address",
    "region",
]
EU_PART_COLUMNS = ["dataset_version_id", "period", "project_no", "row_no", "reg_no"]


class EuProjects:
    source_id = "eu_projects"

    def fetch(self, ctx):
        files = []
        for key in PERIODS:
            files.append(fetch_whole(ctx, f"projects_{key}", f"projects_{key}.csv"))
            files.append(fetch_whole(ctx, f"locations_{key}", f"locations_{key}.csv"))
            files.append(fetch_whole(ctx, f"partners_{key}", f"partners_{key}.csv"))
        for key, (_, schema, *_) in PERIODS.items():
            check_columns(
                f"projects_{key}.csv",
                csv_header(ctx.workdir / f"projects_{key}.csv", delimiter=","),
                schemas()["csv"][schema],
            )
            check_columns(
                f"locations_{key}.csv",
                csv_header(ctx.workdir / f"locations_{key}.csv", delimiter=","),
                schemas()["csv"]["eu_locations"],
            )
            check_columns(
                f"partners_{key}.csv",
                csv_header(ctx.workdir / f"partners_{key}.csv", delimiter=","),
                schemas()["csv"]["eu_partners"],
            )
        return files

    def version_date(self, ctx, files):
        return resource_date(ctx, [f"projects_{k}" for k in PERIODS])

    def load(self, ctx, conn, version_id, files):
        counts: Counter = Counter()
        natural_projects: set[tuple[str, str]] = set()

        def project_rows():
            for key, cols in PERIODS.items():
                period, _schema, start_col, end_col, fund_col, total_col, eu_col = cols
                path = ctx.workdir / f"projects_{key}.csv"
                for line, r in enumerate(iter_csv(path, delimiter=","), start=2):
                    counts["project_rows"] += 1
                    pno = (r["ProjektaNumurs"] or "").strip()
                    if not pno:
                        raise ValueError(f"{path.name} line {line}: empty ProjektaNumurs")
                    atype = clean_text(r["ProjektaIesniedzejaTips"]) or ""
                    natural = atype in NATURAL_TYPES
                    reg = legal_reg_no(r["IesniedzejaRegNo"])
                    if natural:
                        natural_projects.add((period, pno))
                        counts["natural_or_ik"] += 1
                        title = name = None
                        reg = None
                    else:
                        title = clean_text(r["ProjektaNosaukums"])
                        name = clean_text(r["IesniedzejaNosaukums"])
                        if reg is None and (r["IesniedzejaRegNo"] or "").strip():
                            counts["skipped_applicant_reg"] += 1
                    mon = (r["PecuzraudzibasPeriodsGados"] or "").strip()
                    counts["projects_loaded"] += 1
                    yield (
                        version_id,
                        period,
                        pno,
                        title,
                        clean_text(r["ProjektaStatuss"]),
                        atype or None,
                        reg,
                        name,
                        natural,
                        _iso_date(r[start_col], f"{path.name}:{line} start"),
                        _iso_date(r[end_col], f"{path.name}:{line} end"),
                        clean_text(r[fund_col]),
                        clean_text(r["PasakumaNumurs"]),
                        clean_text(r["PasakumaNosaukums"]),
                        _decimal(r[total_col], f"{path.name}:{line} total"),
                        _decimal(r[eu_col], f"{path.name}:{line} eu"),
                        int(mon) if mon.isdigit() else None,
                        _iso_date(
                            r["PedejasDatuAtjauninasanasDatums"], f"{path.name}:{line} upd"
                        ),
                    )

        def location_rows():
            for key, (period, *_rest) in PERIODS.items():
                path = ctx.workdir / f"locations_{key}.csv"
                with_cad = total = 0
                for row_no, r in enumerate(iter_csv(path, delimiter=","), start=1):
                    total += 1
                    pno = (r["ProjektaNumurs"] or "").strip()
                    cad = _cadastre_text(r["KadastraNumurs"])
                    if cad:
                        # 17-digit premise groups → building (14) for matching (D11)
                        if len(cad) == 17 and cad.isdigit():
                            cad = cad[:14]
                        with_cad += 1
                    addr = clean_text(r["Adrese"])
                    if (period, pno) in natural_projects:
                        addr = None  # no address for natural-person / IK projects
                    counts["locations_loaded"] += 1
                    yield (
                        version_id,
                        period,
                        pno,
                        row_no,
                        cad,
                        addr,
                        clean_text(r["Regions"]),
                    )
                counts[f"loc_coverage_{period}"] = {
                    "rows": total,
                    "with_cadastre": with_cad,
                    "share": round(with_cad / total, 4) if total else None,
                }

        def partner_rows():
            for key, (period, *_rest) in PERIODS.items():
                path = ctx.workdir / f"partners_{key}.csv"
                row_no = 0
                for r in iter_csv(path, delimiter=","):
                    counts["partner_rows"] += 1
                    reg = legal_reg_no(r["RegNr"])
                    if reg is None:
                        counts["partners_skipped"] += 1
                        continue
                    row_no += 1
                    counts["partners_loaded"] += 1
                    yield (
                        version_id,
                        period,
                        (r["ProjektaNumurs"] or "").strip(),
                        row_no,
                        reg,
                    )

        n_proj = copy_rows(conn, "ref.eu_project", EU_PROJECT_COLUMNS, project_rows())
        n_loc = copy_rows(conn, "ref.eu_project_location", EU_LOC_COLUMNS, location_rows())
        n_part = copy_rows(conn, "ref.eu_project_partner", EU_PART_COLUMNS, partner_rows())
        ctx.processing = [
            {
                "id": "personal_data",
                "what": "Natural-person/IK projects: no name, title, reg. no., address; "
                "partners: reg. no. only (D11)",
                "natural_or_ik_projects": counts["natural_or_ik"],
                "partners_skipped": counts["partners_skipped"],
            },
            {
                "id": "location_coverage",
                "what": "Share of location rows with a cadastre number (module 12 warning)",
                **{k: v for k, v in counts.items() if k.startswith("loc_coverage_")},
            },
        ]
        return {
            "eu_project": {"source_rows": counts["project_rows"], "loaded": n_proj},
            "eu_project_location": {"loaded": n_loc},
            "eu_project_partner": {"source_rows": counts["partner_rows"], "loaded": n_part},
        }


# --- EIS procurement ----------------------------------------------------------

EIS_COLUMNS = [
    "dataset_version_id",
    "year",
    "row_no",
    "procurement_id",
    "procurement_no",
    "title",
    "buyer_reg_no",
    "buyer_name",
    "winner_reg_no",
    "winner_country",
    "status",
    "amount",
    "currency",
    "signed_on",
    "ended_on",
    "eis_url",
    "iub_url",
]


class EisProcurement:
    source_id = "eis_procurement"

    def fetch(self, ctx):
        files = []
        for role in ctx.cfg.files[self.source_id]:
            year = role.removeprefix("y")
            files.append(fetch_whole(ctx, role, f"eis_{year}.csv"))
            check_columns(
                f"eis_{year}.csv",
                csv_header(ctx.workdir / f"eis_{year}.csv", delimiter=","),
                schemas()["csv"]["eis_results"],
            )
        return files

    def version_date(self, ctx, files):
        return resource_date(ctx, list(ctx.cfg.files[self.source_id]))

    def load(self, ctx, conn, version_id, files):
        counts: Counter = Counter()

        def rows():
            for role in ctx.cfg.files[self.source_id]:
                year = int(role.removeprefix("y"))
                path = ctx.workdir / f"eis_{year}.csv"
                for line, r in enumerate(iter_csv(path, delimiter=","), start=2):
                    counts["source_rows"] += 1
                    winner = legal_reg_no(r["Uzvaretaja_registracijas_numurs"])
                    if not winner and (r["Uzvaretaja_registracijas_numurs"] or "").strip():
                        counts["winner_reg_skipped"] += 1
                    # Uzvaretaja_nosaukums never stored (D11)
                    counts["loaded"] += 1
                    yield (
                        version_id,
                        year,
                        counts["loaded"],
                        clean_text(r["Iepirkuma_ID"]),
                        clean_text(r["Iepirkuma_identifikacijas_numurs"]),
                        clean_text(r["Iepirkuma_nosaukums"]),
                        legal_reg_no(r["Pasutitaja_registracijas_numurs"]),
                        clean_text(r["Pasutitaja_nosaukums"]),
                        winner,
                        clean_text(r["Uzvaretaja_valsts"]),
                        clean_text(r["Iepirkuma_statuss"]),
                        _decimal(r["Aktuala_liguma_summa"], f"{path.name}:{line}"),
                        clean_text(r["Aktuala_liguma_summas_valuta"]),
                        _iso_date(r["Liguma_dok_noslegsanas_datums"], f"{path.name}:{line} signed"),
                        _iso_date(r["Izbeigsanas_datums"], f"{path.name}:{line} ended"),
                        clean_text(r["Hipersaite_EIS_kura_pieejams_zinojums"]),
                        clean_text(r["Hipersaite_uz_IUB_publikaciju"]),
                    )

        loaded = copy_rows(conn, "ref.eis_contract", EIS_COLUMNS, rows())
        ctx.processing = [
            {
                "id": "personal_data",
                "what": "Winner name never stored; winner reg. no. legal-entity pattern only (D11)",
                "winner_reg_skipped": counts["winner_reg_skipped"],
            },
            {
                "id": "coverage",
                "what": "Only procurements run in the EIS e-konkursi subsystem",
            },
        ]
        return {"eis_contract": {"source_rows": counts["source_rows"], "loaded": loaded}}


# --- FID frozen real estate ---------------------------------------------------

FID_FROZEN_COLUMNS = [
    "dataset_version_id",
    "row_no",
    "cadastre_nr",
    "asset_type",
    "subject_link",
    "holder_reg_no",
    "mark_ta",
    "mark_vzd",
    "basis",
    "sanctions_from",
    "legal_act",
    "fid_checked",
    "published_on",
]


def _cell(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return clean_text(str(value))


class FidFrozenProperty:
    source_id = "fid_frozen_property"
    FILE = "fid_ta_vzd.xlsx"

    def _sheet(self, ctx):
        spec = schemas()["xlsx"]["fid_ta_vzd"]
        wb = load_workbook(ctx.workdir / self.FILE, read_only=True, data_only=True)
        if spec["sheet"] not in wb.sheetnames:
            wb.close()
            raise SchemaMismatch(f"{self.FILE}: no sheet {spec['sheet']!r}")
        return wb, wb[spec["sheet"]], spec

    def check_schema(self, ctx):
        wb, ws, spec = self._sheet(ctx)
        try:
            header_row = int(spec.get("header_row", 1))
            header = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
            titles = [str(c).strip() if c is not None else "" for c in header]
            while titles and titles[-1] == "":
                titles.pop()
            check_columns(self.FILE, titles, spec["columns"])
        finally:
            wb.close()

    def fetch(self, ctx):
        files = [fetch_whole(ctx, "ta_vzd", self.FILE)]
        self.check_schema(ctx)
        return files

    def version_date(self, ctx, files):
        return resource_date(ctx, ["ta_vzd"])

    def load(self, ctx, conn, version_id, files):
        self.check_schema(ctx)
        counts: Counter = Counter()
        wb, ws, spec = self._sheet(ctx)
        start = int(spec.get("header_row", 1)) + 1

        def rows():
            for row_no, cells in enumerate(ws.iter_rows(min_row=start, values_only=True), start=1):
                cells = list(cells) + [None] * 18
                if all(c is None or str(c).strip() == "" for c in cells[:18]):
                    continue
                counts["source_rows"] += 1
                cad = _cadastre_text(cells[5])
                if not cad:
                    counts["skipped_no_cadastre"] += 1
                    continue
                # Never store cols 1, 3, 7, 13, 14; col 2 only as legal-entity reg. no.
                holder = legal_reg_no(cells[1])
                pub = cells[17]
                if isinstance(pub, datetime):
                    published = pub.date()
                elif isinstance(pub, date):
                    published = pub
                else:
                    published = _iso_date(_cell(pub), f"row {row_no} published") if pub else None
                counts["loaded"] += 1
                yield (
                    version_id,
                    row_no,
                    cad,
                    _cell(cells[3]),
                    _cell(cells[4]),
                    holder,
                    _cell(cells[9]),
                    _cell(cells[10]),
                    _cell(cells[11]),
                    _cell(cells[14]),
                    _cell(cells[15]),
                    _cell(cells[16]),
                    published,
                )

        try:
            loaded = copy_rows(conn, "ref.fid_frozen_property", FID_FROZEN_COLUMNS, rows())
        finally:
            wb.close()
        ctx.processing = [
            {
                "id": "personal_data",
                "what": "No subject name, birth date, address or sanctioned-person columns; "
                "holder reg. no. legal-entity only (D11)",
                "skipped_no_cadastre": counts["skipped_no_cadastre"],
            }
        ]
        return {"fid_frozen_property": {"source_rows": counts["source_rows"], "loaded": loaded}}
