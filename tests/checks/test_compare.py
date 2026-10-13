"""The test-set comparison tool on a filled copy of tests/data/test_set_template.xlsx."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.batch.headers import canonical
from engine.checks import compare

pytestmark = pytest.mark.stack

TEMPLATE = Path(__file__).resolve().parents[1] / "data" / "test_set_template.xlsx"
HINT = "‹Ir: … / Nav / Nav datu / Nav pārklājuma›"
# projekta_nr, vertiba, m1..m6, parbaudes_datums (the template's example rows stay in)
ROWS = [
    (
        "P-1",
        9000050400,
        "Ir: zemes vienība",
        HINT,
        "Nav",
        "Ir: plūdu zonā",
        "Nav",
        "nav",
        date(2026, 10, 1),
    ),
    ("P-1", "05000010069", "Ir", None, "Nav", "Nav", None, None, date(2026, 1, 1)),
    (
        "P-2",
        "Institūta iela 4, Jelgava",
        "Ir",
        "Nav (atbilst)",
        "Ir: dabas parks",
        "Nav pārklājuma",
        "Ir: BIS-A",
        "Ir",
        date(2026, 10, 1),
    ),
    ("P-3", "12345", "Ir", None, None, None, None, None, None),
    # property-or-parcel number; veids picks the parcel; its application names another
    # object's address (Salnas iela 22 belongs to 09000050405): M2 "Neatbilst"
    ("P-4", "05000040014", "Ir", "Ir: neatbilst", None, None, None, None, None),
    # the application's address is the object's own: M2 agrees "Nav"
    ("P-5", "09000360220", "Ir", "Nav (atbilst)", None, None, None, None, None),
]
KIND = {8: "zemes vienība"}  # row -> veids
ADDRESS = {  # row -> "Adrese pieteikumā"
    5: "Neesoša iela 1, Jelgava",  # not found: the object is still checked, M2 not compared
    8: "Salnas iela 22, Jelgava",
    9: "Institūta iela 4, Jelgava",
}


@pytest.fixture
def filled(tmp_path):
    wb = load_workbook(TEMPLATE)
    ws = wb[compare.SHEET]
    col = {canonical(c.value): c.column for c in ws[1]}  # by column, not position or header text
    for r, (project, value, *answers, checked) in enumerate(ROWS, start=4):
        ws.cell(row=r, column=col["projekta_nr"], value=project)
        ws.cell(row=r, column=col["vertiba"], value=value)  # 9000050400 as a number: zero lost
        ws.cell(row=r, column=col["veids"], value=KIND.get(r))
        ws.cell(row=r, column=col["adrese"], value=ADDRESS.get(r))
        for name, a in zip(compare.MODULE_COLUMNS, answers, strict=True):
            ws.cell(row=r, column=col[name], value=a)
        ws.cell(row=r, column=col["parbaudes_datums"], value=checked)
    out = tmp_path / "filled.xlsx"
    wb.save(out)
    return out


@pytest.mark.parametrize(
    ("text", "status"),
    [
        ("Ir: mikroliegums", "found"),
        ("Nav", "clear"),
        ("nav datu", "no_data"),
        ("Nav pārklājuma – nav kartes", "not_covered"),
        ("NAV PARKLAJUMA", "not_covered"),
        (HINT, None),
        ("", None),
        (None, None),
        ("varbūt", None),
    ],
)
def test_parse_answer(text, status):
    assert compare.parse_answer(text) == status


def test_compare_filled_template(engine_db, filled, tmp_path):
    with engine_db() as conn:
        report = compare.compare(conn, filled)
    lines = {(line.row_no, line.module): line for line in report.lines}
    assert {line.row_no for line in report.lines} == {4, 5, 6, 7, 8, 9}  # examples skipped
    assert report.rule_set == "3 (agreed)"

    flood = lines[(4, "flood_10y")]
    assert (flood.designation, flood.manual, flood.engine, flood.outcome) == (
        "09000050400",
        "found",
        "found",
        "agrees",
    )
    assert lines[(4, "address_consistency")].outcome == "not_answered"  # hint left in
    assert lines[(4, "encumbrances")].outcome == "agrees"

    differs = lines[(5, "flood_10y")]
    assert (differs.manual, differs.engine, differs.outcome) == ("clear", "not_covered", "differs")
    assert differs.newer_data is True  # data dated after the 2026-01-01 check
    assert lines[(5, "address_consistency")].outcome == "not_answered"
    unresolved = lines[(5, "address_consistency")]  # the address in the application: not found
    assert unresolved.engine == "not_applicable" and "norādītā adrese" in unresolved.note
    assert lines[(5, "flood_10y")].designation == "05000010069"  # the object is still checked

    # "Adrese pieteikumā" is the given address: M2 can be compared
    mismatch = lines[(8, "address_consistency")]
    assert (mismatch.manual, mismatch.engine, mismatch.outcome) == ("found", "found", "agrees")
    assert mismatch.detail == "Neatbilst"
    own = lines[(9, "address_consistency")]
    assert (own.manual, own.engine, own.outcome) == ("clear", "clear", "agrees")

    addr = lines[(6, "address_consistency")]
    assert (addr.designation, addr.engine, addr.outcome) == ("09000360220", "clear", "agrees")
    assert "2 objekti" in addr.note
    assert lines[(6, "nature")].outcome == "agrees"
    assert lines[(6, "flood_10y")].outcome == "agrees"
    assert lines[(6, "bis")].detail.startswith("BIS-A (building_on_parcel)")

    bad = lines[(7, "cadastre_facts")]
    assert (bad.designation, bad.engine, bad.outcome) == (None, None, "not_comparable")

    picked = lines[(8, "cadastre_facts")]  # veids: no "neviennozīmīgs" note
    assert (picked.designation, picked.engine, picked.note) == ("05000040014", "found", None)
    assert "Zemes vienība" in picked.detail

    counts = report.counts()
    assert counts["flood_10y"] == {
        "agrees": 2,
        "differs": 1,
        "not_answered": 3,  # rows 7 (not resolved), 8 and 9 (only m1, m2 answered)
        "not_comparable": 0,
    }
    assert counts["cadastre_facts"]["not_comparable"] == 1  # row 7 not resolved

    out = tmp_path / "report.xlsx"
    compare.write_xlsx(report, out)
    ws = load_workbook(out)["Salidzinajums"]
    assert ws.cell(row=2, column=4).value == "09000050400"  # text, leading zero kept
    assert ws.cell(row=2, column=4).number_format == "@"
