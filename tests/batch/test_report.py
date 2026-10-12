"""Excel report (PLAN.md §7.1): round trip through openpyxl, codes as text,
T-E4 (Excel half), run vs replay, unresolved rows, the test-set template as
input, and a 1,000-row batch."""

import io
import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from engine.batch.headers import canonical
from engine.batch.plan import run_batch
from engine.batch.reader import read_upload
from engine.checks import runner
from engine.checks.modules import MODULES
from reports import excel
from tests.batch.conftest import STORED_NAME, TE4_BUILDING

pytestmark = pytest.mark.stack

TEMPLATE = Path(__file__).resolve().parents[1] / "data" / "test_set_template.xlsx"
SHEETS = ["Kopsavilkums", "Pieteicēji", "Detaļas", "Avoti", "Neatrasts"]
HEADERS = {
    "Kopsavilkums": [
        "Nr.",
        "Rinda",
        "Projekta Nr.",
        "Ievadītā vērtība",
        "Pārbaudītais objekts",
        "Objekta veids",
        "Interpretācija",
        "Adrese (kadastrā)",
        "Norādītā adrese",
        "Pieteicēja reģ. nr.",
        # object modules only; the applicant modules have their own sheet
        *[f"{m.no}. {m.name_lv}" for m in MODULES if m.subject == "object"],
        "Kopējais risks",
        "Brīdinājumi",
    ],
    "Detaļas": [
        "Nr.",
        "Pārbaudītais objekts",
        "Vienība",
        "Modulis",
        "Statuss",
        "Ieraksts",
        "Pārklājums, m²",
        "Pārklājums, %",
        "Atbilstības veids",
        "Risks",
        "Noteikums",
        "Noteikuma teksts",
        "Avots",
        "Datu datums",
        "data.gov.lv saite",
        "Avots sistēmā",
    ],
    "Pieteicēji": [
        "Pieteicēja reģ. nr.",
        "Nosaukums",
        "Forma",
        "Statuss UR",
        "Reģistrēts",
        "Darbību beidzis",
        "Juridiskā adrese",
        "Atvērti procesi",
        "Objekti (Nr.)",
        *[f"{m.no}. {m.name_lv}" for m in MODULES if m.subject == "applicant"],
        "Kopējais risks",
        "Brīdinājumi",
    ],
    "Avoti": [
        "Avots",
        "Avota id",
        "Datu kopas versija",
        "Datu datums",
        "Svaigums",
        "Svaiguma pamatojums",
        "Avots atjaunots",
        "Brīdināt, ja avots nav atjaunots ilgāk par",
        "Pārklājums",
        "Apstrādes piezīmes",
        "data.gov.lv saite",
    ],
    "Neatrasts": [
        "Rinda",
        "Kolonna",
        "Ievadīts",
        "Projekta Nr.",
        "Statuss",
        "Iemesls",
        "Ieteikumi",
    ],
}
# vertiba, projekta_nr, adrese, pieteicejs_reg_nr, veids
INPUT = [
    ("vertiba", "projekta_nr", "adrese", "pieteicejs_reg_nr", "veids"),
    ("09000360220", "007", "Institūta iela 4, Jelgava", "40003000000", None),
    (9000050400, "007", None, None, None),  # a number: leading zero lost in Excel
    ("05000040014", "008", None, None, None),  # property or parcel: both
    (TE4_BUILDING, "009", None, None, None),
    ("13000102702", "010", "Neesoša iela 1, Jelgava", "010190-12345", None),
    ("123", "011", None, None, None),
    ("=1+1", "=HYPERLINK(1)", None, None, None),  # never a formula in the report
]


def xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(list(r))
    for row in wb.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # typed as text, as a user's cell would be
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def open_report(data) -> object:
    return load_workbook(io.BytesIO(excel.to_bytes(data)))


def values(wb) -> dict[str, list[tuple]]:
    return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def column(ws, name):
    return [c.value for c in ws[1]].index(name) + 1


@pytest.fixture(scope="module")
def run(batch_db):
    with batch_db() as c:
        return run_batch(c, read_upload("saraksts.xlsx", xlsx(INPUT), 1000))


@pytest.fixture(scope="module")
def wb(batch_db, run):
    with batch_db() as c:
        return open_report(excel.load(c, run))


def test_sheets_headers_frozen_rows_filters_and_footer(wb):
    assert wb.sheetnames == SHEETS
    for name in SHEETS:
        ws = wb[name]
        assert [c.value for c in ws[1]] == HEADERS[name]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref.startswith("A1:")
        assert ws.oddFooter.center.text == excel.FOOTER
        texts = [c.value for c in ws["A"] if isinstance(c.value, str)]
        assert excel.FOOTER in texts


def test_codes_are_text_with_leading_zeros(wb):
    ws = wb["Kopsavilkums"]
    obj, proj, inp = (
        column(ws, "Pārbaudītais objekts"),
        column(ws, "Projekta Nr."),
        column(ws, "Ievadītā vērtība"),
    )
    by_object = {ws.cell(r, obj).value: r for r in range(2, ws.max_row + 1)}
    r = by_object["09000050400"]
    cell = ws.cell(r, obj)
    assert cell.data_type == "s" and cell.number_format == "@"
    assert ws.cell(r, proj).value == "007" and ws.cell(r, proj).data_type == "s"
    assert ws.cell(r, inp).value == "9000050400"  # as typed (the number)
    reg = ws.cell(by_object["09000360220"], column(ws, "Pieteicēja reģ. nr.")).value
    assert reg == "40003000000"
    d = wb["Detaļas"]
    units = [d.cell(i, 3) for i in range(2, d.max_row + 1) if d.cell(i, 3).value]
    assert all(u.data_type == "s" and u.value[0].isdigit() for u in units)
    assert any(u.value.startswith("0") for u in units)


def test_t_e4_name_round_trips_through_excel(wb):
    """T-E4, Excel half: a name read with CR, LF, tab and a backslash is stored
    normalised (Phase 2) and comes back from the report unchanged."""
    d = wb["Detaļas"]
    texts = [d.cell(r, 6).value for r in range(2, d.max_row + 1)]
    line = next(t for t in texts if t and TE4_BUILDING in t and "Būve" in t)
    assert STORED_NAME in line and "\\" in line
    assert not any(ch in line for ch in "\r\n\t")


def test_module_cells_risk_colours_and_labels(wb):
    ws = wb["Kopsavilkums"]
    m2 = column(ws, f"2. {MODULES[1].name_lv}")
    risk = column(ws, "Kopējais risks")
    label = column(ws, "Interpretācija")
    rows = {(ws.cell(r, 5).value, ws.cell(r, label).value): r for r in range(2, ws.max_row + 1)}
    assert ("05000040014", "divdomīgs: nekustamais īpašums") in rows
    assert ("05000040014", "divdomīgs: zemes vienība") in rows
    assert ws.cell(rows[("09000050400", None)], m2).value == "nav piemērojams"
    assert ws.cell(rows[("13000102702", None)], m2).value == excel.ADDRESS_UNRESOLVED
    assert ws.cell(rows[("09000360220", None)], m2).value == "Nav / nav"  # address matches
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, risk)
        if cell.value is None:
            continue
        level = {v: k for k, v in excel.RISK_LV.items()}[cell.value]
        assert cell.fill.fgColor.rgb.endswith(excel.RISK_FILL[level])


def test_details_name_rules_sources_and_links(wb):
    d = wb["Detaļas"]
    rows = [[c.value for c in r] for r in d.iter_rows(min_row=2) if isinstance(r[0].value, int)]
    flood = next(r for r in rows if r[1] == "09000050400" and r[3].startswith("4."))
    assert flood[6] == pytest.approx(100.0) and flood[7] == pytest.approx(1.0)
    assert flood[10] == "flood_10y_major" and flood[11]
    assert flood[13] == "2026-10-10" and flood[14].startswith("https://data.gov.lv/")
    bis = [r for r in rows if r[3].startswith("5.") and r[1] == "09000360220"]
    assert any("BIS-A" in r[5] for r in bis)
    frozen = [r for r in rows if r[1] == "09000360220" and "2025-01-31" in (r[5] or "")]
    assert frozen and all("7399000000" in r[5] for r in frozen)
    # "Avots sistēmā": the item in the publisher's system (engine/links.py)
    assert any("BIS-A" in r[5] and r[15] == "BIS" for r in bis)
    case_cell = next(
        d.cell(i, 16) for i in range(2, d.max_row + 1)
        if "BIS-A" in str(d.cell(i, 6).value) and d.cell(i, 2).value == "09000360220"
    )  # fmt: skip
    assert "planned_constructions?search%5Bcase_number%5D=BIS-A" in case_cell.hyperlink.target
    nature = next(r for r in rows if r[1] == "09000360220" and r[3].startswith("3."))
    assert nature[15] == f"Ozols (DAP) (meklēt: {nature[2]})"  # no deep link: search + number
    technical = [r for r in rows if r[1] == "09000050400" and r[3].startswith("6.")]
    assert len(technical) == 1 and "Tehniskie apgrūtinājumi" in technical[0][5]
    assert technical[0][9] == "nav" and "meklēt: 8317010100" in technical[0][15]
    # a property's lines come from its units' own results: own rules, own risk
    facts = {r[2]: r for r in rows if r[1] == "09000360220" and r[3].startswith("1.")}
    assert facts["09000360220"][10] == "cadastre_facts_found"
    assert "cadastre_prereg" in facts["09000360220012"][10]  # the pre-registered building
    gaps = [r for r in rows if r[4] == "Nav pārklājuma"]
    assert gaps and all(r[5] for r in gaps)  # a data gap always says why


def test_sources_sheet_has_versions_rule_set_and_thresholds(wb, batch_db, run):
    ws = wb["Avoti"]
    cells = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    with batch_db() as c:
        sha = c.execute("SELECT rule_set_sha256 FROM work.run WHERE id = %s", (run,)).fetchone()[0]
    assert sha in cells and "melnraksts" in cells  # rule set v3 is a draft
    assert any(isinstance(v, str) and "1 cm" in v for v in cells)  # flood processing note
    assert "Visa valsts (tikai kartētie biotopi)" in cells  # coverage in Latvian
    assert not any(isinstance(v, str) and v.startswith("Whole country") for v in cells)
    assert excel.THRESHOLD_LV["min_overlap_m2"] in cells
    ids = [ws.cell(r, 2).value for r in range(2, 20)]
    assert "flood_10y" in ids and "cadastre_text" in ids


def test_unresolved_rows_land_in_neatrasts(wb):
    ws = wb["Neatrasts"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2) if isinstance(r[0].value, int)]
    by = {(r[0], r[1]): r for r in rows}
    assert by[(7, "vērtība")][2] == "123" and by[(7, "vērtība")][4] == "nederīgs"
    assert by[(6, "adrese")][2] == "Neesoša iela 1, Jelgava" and by[(6, "adrese")][5]
    reg = by[(6, "pieteicēja reģ. nr.")]
    assert reg[2] is None and "010190" not in str(rows)  # personal code never shown
    formula = ws.cell([r[0] for r in rows].index(8) + 2, 3)
    assert formula.value == "=1+1" and formula.data_type == "s"


def test_report_from_replay_equals_report_from_run(batch_db, run):
    with batch_db() as c:
        stored = values(open_report(excel.load(c, run)))
        again = values(open_report(excel.load(c, run)))
        replayed = values(open_report(excel.load(c, run, rows=runner.replay(c, run))))
    assert stored == again
    assert stored == replayed


def test_test_set_template_as_input(conn, tmp_path):
    wb = load_workbook(TEMPLATE)
    ws = wb["Testa_kopa"]
    col = {canonical(c.value): c.column for c in ws[1]}
    for r, (value, kind) in enumerate(
        [("05000040014", "zemes vienība"), (9000050400, None), ("123", None)], start=4
    ):
        ws.cell(r, col["projekta_nr"], "T-1")
        ws.cell(r, col["vertiba"], value)
        ws.cell(r, col["veids"], kind)
        ws.cell(r, col["parbaudes_datums"], date(2026, 10, 1))
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    rows = read_upload(path.name, path.read_bytes(), 1000)
    assert [r.row_no for r in rows] == [4, 5, 6]  # the two examples skipped
    rid = run_batch(conn, rows)
    report = open_report(excel.load(conn, rid))
    k = report["Kopsavilkums"]
    checked = [(k.cell(r, 5).value, k.cell(r, 6).value) for r in range(2, k.max_row + 1)]
    assert ("05000040014", "Zemes vienība") in checked and (
        "09000050400",
        "Zemes vienība",
    ) in checked
    assert len([c for c in checked if c[0]]) == 2
    assert report["Neatrasts"].cell(2, 3).value == "123"


def test_thousand_row_batch(conn):
    """1,000 input rows (pilot target < 5 min, PLAN §9). The fixture repeats a few
    objects; the real timing is in tests/pilot/test_pilot_batch.py."""
    values_ = ["09000360220", "09000050400", "13000102702", "05000040014", "123"]
    text = "vertiba\n" + "\n".join(values_[i % len(values_)] for i in range(1000))
    t0 = time.perf_counter()
    rows = read_upload("saraksts.csv", text.encode(), 1000)
    rid = run_batch(conn, rows)
    data = excel.to_bytes(excel.load(conn, rid))
    elapsed = time.perf_counter() - t0
    wb = load_workbook(io.BytesIO(data), read_only=True)
    assert wb["Kopsavilkums"].max_row >= 1000  # 800 rows + 200 checked both ways + footer
    assert len(runner.stored_inputs(conn, rid)) == 1000
    print(f"1,000 rows (fixture): {elapsed:.1f} s")
    assert elapsed < 120


@pytest.mark.parametrize(
    ("text", "lv"),
    [("14 days", "14 d."), ("1 year 1 mon", "1 g. 1 mēn."), ("2 days", "2 d."), (None, None)],
)
def test_max_age_in_latvian(text, lv):
    assert excel._age(text) == lv
