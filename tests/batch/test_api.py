"""POST /api/batch, the report download, the web page and the CLI (Phase 5)."""

import io
import re

import pytest
from openpyxl import load_workbook

from engine.batch.headers import INPUT_COLUMNS
from engine.checks.__main__ import main
from tests.batch.test_report import SHEETS, values, xlsx
from tests.etl.conftest import ETL_DB

pytestmark = pytest.mark.stack

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_pasted_list_gives_a_run_and_a_report(client):
    r = client.post("/api/batch", data={"text": "09000050400\n123\n05000040014"})
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["origin"] == "batch" and body["targets"] == 3  # 05000040014 twice
    assert body["inputs"] == {"rows": 3, "checked": 2, "not_checked": 1}
    rep = client.get(body["report_url"])
    assert rep.status_code == 200 and rep.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(rep.content))
    assert wb.sheetnames == SHEETS
    assert wb["Neatrasts"].cell(2, 3).value == "123"


def test_file_upload_returns_the_excel_directly(client):
    data = xlsx([("vertiba", "projekta_nr"), (9000050400, "007")])
    r = client.post(
        "/api/batch",
        files={"file": ("saraksts.xlsx", data, XLSX)},
        data={"format": "xlsx", "text": "ignored when a file is given"},
    )
    assert r.status_code == 200  # a download, not the JSON "created" answer
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"].startswith('attachment; filename="kadastra_parbaude_')
    ws = load_workbook(io.BytesIO(r.content))["Kopsavilkums"]
    assert (ws.cell(2, 3).value, ws.cell(2, 5).value) == ("007", "09000050400")
    assert ws.max_row == 4  # one object + blank + footer


def test_bad_input_is_refused_with_a_reason(client):
    r = client.post("/api/batch", data={"text": "  "})
    assert r.status_code == 422 and "nav ievadīts" in r.json()["detail"]
    r = client.post("/api/batch", data={"text": "09000050400\n" * 1001})
    assert r.status_code == 422 and "1000" in r.json()["detail"]
    r = client.post("/api/batch", data={"text": "09000050400", "sources": "nope"})
    assert r.status_code == 422
    page = client.post("/api/batch", data={"text": "", "format": "xlsx"})
    assert page.status_code == 422 and "text/html" in page.headers["content-type"]
    assert "nav ievadīts" in page.text and "<script" not in page.text


def test_unknown_run_report_is_404(client):
    assert client.get("/api/runs/999999/report.xlsx").status_code == 404
    assert client.get("/api/runs/999999/report.html").status_code == 404
    assert client.get("/api/runs/999999/report.pdf").status_code == 404


def test_html_and_pdf_reports(client):
    r = client.post("/api/batch", data={"text": "09000050400"})
    assert r.status_code == 201, r.text
    run_id = r.json()["id"]
    html = client.get(f"/api/runs/{run_id}/report.html")
    assert html.status_code == 200
    assert "text/html" in html.headers["content-type"]
    assert "Kadastra pārbaude" in html.text
    assert "Informatīvs pārskats" in html.text
    pdf = client.get(f"/api/runs/{run_id}/report.pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")


def test_input_page_and_template(client):
    page = client.get("/")  # the results view (web/, built in the image)
    assert page.status_code == 200 and '<div id="root">' in page.text
    script = re.search(r'src="(/assets/[^"]+\.js)"', page.text).group(1)
    assert client.get(script).status_code == 200
    t = client.get("/api/batch/template.xlsx")
    wb = load_workbook(io.BytesIO(t.content))
    assert wb.sheetnames == ["Saraksts", "Instrukcija"]
    ws = wb["Saraksts"]
    assert [c.value for c in ws[1]] == [c.header for c in INPUT_COLUMNS]
    assert ws["A1"].fill.fgColor.rgb.endswith("1F3A5F")  # required: dark
    assert ws["B1"].fill.fgColor.rgb.endswith("D9E2EF")  # optional: light
    assert ws["A1"].comment.text.startswith("Obligāta")
    assert ws["A2"].number_format == "@" and ws["C2"].number_format == "@"
    (kinds,) = ws.data_validations.dataValidation
    assert str(kinds.sqref).startswith("E2") and "zemes vienība" in kinds.formula1


def test_cli_batch_and_report_replay(batch_db, tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("DATABASE_URL", ETL_DB)
    src = tmp_path / "saraksts.csv"
    src.write_text("vertiba;projekta_nr\n09000360220;P-1\n123;P-2\n", encoding="utf-8")
    assert main(["batch", str(src)]) == 0
    out = capsys.readouterr().out
    assert "2 rows, 1 not checked" in out
    report = next(tmp_path.glob("saraksts_parbaude_*.xlsx"))
    run_id = int(report.stem.rsplit("_", 1)[1])
    replayed = tmp_path / "replay.xlsx"
    assert main(["report", str(run_id), "--replay", "--out", str(replayed)]) == 0
    assert values(load_workbook(report)) == values(load_workbook(replayed))
    assert main(["report", "999999"]) == 1
