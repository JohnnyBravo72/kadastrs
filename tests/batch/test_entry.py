"""Phase 7 API: entry dialog rows (`entries`), the import preview, the source
registry, Žurnāls, new runs from a stored run and project groups (PLAN.md §8.1,
§8.2)."""

import io
import json
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from tests.batch.test_report import xlsx

pytestmark = pytest.mark.stack

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TAUTAS_11A = "Tautas iela 11A, Daugavpils"  # parcel 05000040013 and building 05000041006001


def journal_ids(client, **params) -> list[int]:
    """Run ids on one Žurnāls page (GET /api/runs)."""
    page = client.get("/api/runs", params={"limit": 200, **params}).json()
    return [x["id"] for x in page["runs"]]


def run_count(conn) -> int:
    return conn.execute("SELECT count(*) FROM work.run").fetchone()[0]


def targets(client, run_id):
    return client.get(f"/api/runs/{run_id}/targets").json()


# --- entry dialog rows -----------------------------------------------------------------------


def test_entries_check_only_the_ticked_units(client, conn):
    units = client.get("/api/property/09000360220/units").json()["units"]
    ticked = [{"type": u["type"], "designation": u["designation"]} for u in units[:1]]
    entries = [
        {"value": "09000360220", "kind": "property", "project_no": "P-7", "selected": ticked},
        {"value": TAUTAS_11A, "selected": [{"type": "parcel", "designation": "05000040013"}]},
    ]
    r = client.post("/api/batch", data={"entries": json.dumps(entries)})
    assert r.status_code == 201, r.text
    t = targets(client, r.json()["id"])
    assert [(x["type"], x["designation"], x["project_no"]) for x in t] == [
        ("property", "09000360220", "P-7"),
        ("parcel", "05000040013", None),  # the building at the address was not ticked
    ]
    assert [u["designation"] for u in t[0]["units"]] == [ticked[0]["designation"]]
    detail = client.get(f"/api/runs/{r.json()['id']}/targets/0").json()
    nature = next(m for m in detail["results"] if m["module"] == "nature")["result"]
    assert any(
        w["code"] == "units_selected" and w["message"].startswith(f"pārbaudītas 1 no {len(units)} ")
        for w in nature["warnings"]
    )


def test_entry_kind_decides_an_ambiguous_number(client):
    # 05000040014 is a parcel of 05000041006 and also a property number
    entries = [{"value": "05000040014", "kind": "parcel"}]
    r = client.post("/api/batch", data={"entries": json.dumps(entries)})
    assert [(x["type"], x["label"]) for x in targets(client, r.json()["id"])] == [("parcel", None)]


def test_bad_entries_are_refused(client):
    r = client.post("/api/batch", data={"entries": json.dumps([{"value": ""}])})
    assert r.status_code == 422 and "ievades loga" in r.json()["detail"]
    r = client.post("/api/batch", data={"entries": "[]"})
    assert r.status_code == 422 and "nav ievadīts" in r.json()["detail"]


# --- import preview ----------------------------------------------------------------------------


PREVIEW_FILE = [
    ("vertiba", "projekta_nr", "pieteicejs_reg_nr"),
    (9000050400, "P-1", "40003000000"),  # leading zero lost in Excel
    ("09000050400", "P-1", None),  # duplicate of the row above
    ("123", "P-2", None),  # invalid
    ("010190-12345", None, None),  # looks like a personal code
    ("05000040014", None, "010190-12345"),  # ambiguous; reg. no. not a legal entity
]


def test_preview_shows_types_corrections_duplicates_and_stores_nothing(client, conn):
    before = run_count(conn)
    r = client.post(
        "/api/batch/preview", files={"file": ("saraksts.xlsx", xlsx(PREVIEW_FILE), XLSX)}
    )
    assert r.status_code == 200, r.text
    rows = {x["row_no"]: x for x in r.json()["rows"]}
    assert run_count(conn) == before
    assert rows[2]["type"] == "parcel" and rows[2]["status"] == "resolved" and rows[2]["checked"]
    assert rows[2]["corrections"] and rows[2]["applicant_reg_no"] == "40003000000"
    assert rows[3]["duplicate_of"] == 2 and any("atkārtojas" in w for w in rows[3]["warnings"])
    assert rows[4]["status"] == "invalid" and not rows[4]["checked"]
    # a personal code is neither shown nor kept
    assert rows[5]["value"] is None and rows[5]["status"] == "not_stored"
    assert "010190" not in r.text
    assert [t["label"] for t in rows[6]["targets"]] == [
        "divdomīgs: nekustamais īpašums",
        "divdomīgs: zemes vienība",
    ]
    assert any(p["column"] == "pieteicejs_reg_nr" for p in rows[6]["problems"])


def test_rows_removed_in_the_preview_are_not_checked(client):
    data = xlsx(PREVIEW_FILE)
    r = client.post(
        "/api/batch",
        files={"file": ("saraksts.xlsx", data, XLSX)},
        data={"skip_rows": "3,4"},
    )
    assert r.status_code == 201, r.text
    inputs = client.get(f"/api/runs/{r.json()['id']}/inputs").json()
    assert [i["row_no"] for i in inputs] == [2, 5, 6]
    report = load_workbook(io.BytesIO(client.get(r.json()["report_url"]).content))
    not_found = [row[0] for row in report["Neatrasts"].iter_rows(min_row=2, values_only=True)]
    assert 4 not in not_found  # the invalid row was removed, not reported


def test_preview_refuses_bad_input(client):
    r = client.post("/api/batch/preview", data={"text": "  "})
    assert r.status_code == 422
    r = client.post("/api/batch", data={"text": "09000050400", "skip_rows": "x"})
    assert r.status_code == 422


# --- sources -----------------------------------------------------------------------------------


def test_source_registry_for_the_checklist(client):
    r = client.get("/api/sources")
    assert r.status_code == 200
    by = {s["id"]: s for s in r.json()}
    assert by["flood_10y"]["selectable"] and by["flood_10y"]["partial_coverage"]
    assert by["nature_habitats"]["partial_coverage"]
    assert not by["nature_microreserves"]["partial_coverage"]
    assert [m["id"] for m in by["bis_cases"]["modules"]] == ["bis"]
    # base data every module reads cannot be unticked
    assert not by["cadastre_spatial"]["selectable"] and by["cadastre_spatial"]["used_by"]
    assert by["cadastre_text"]["selectable"]  # selects module 1
    assert by["flood_10y"]["live_version"]["dataset_date"]
    assert by["flood_10y"]["datagov_urls"][0].startswith("https://data.gov.lv/dati/lv/dataset/")
    assert by["flood_10y"]["group"] == "Plūdi"


# --- Žurnāls, new runs, groups -------------------------------------------------------------------


@pytest.fixture(scope="module")
def first(client):
    r = client.post(
        "/api/batch", data={"text": "vertiba;projekta_nr\n09000050400;P-1\n13000102702;P-2"}
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_journal_lists_runs_and_reopens_them(client, first):
    page = client.get("/api/runs", params={"limit": 200}).json()
    runs = page["runs"]
    assert page["total"] >= len(runs) and page["limit"] == 200 and page["offset"] == 0
    row = next(x for x in runs if x["id"] == first)
    assert row["targets"] == 2 and row["projects"] == ["P-1", "P-2"]
    assert row["risk"] == "high"  # 09000050400 is in the flood finding
    assert [x["id"] for x in runs] == sorted((x["id"] for x in runs), reverse=True)
    assert client.get(f"/api/runs/{first}").json()["id"] == first


def test_api_rerun_both_kinds(client, first):
    sources = [s for s in client.get(f"/api/runs/{first}").json()["sources"] if s != "flood_10y"]
    r = client.post(f"/api/runs/{first}/rerun", json={"kind": "sources", "sources": sources})
    assert r.status_code == 201, r.text
    child = r.json()
    assert child["parent_run_id"] == first and child["rerun_kind"] == "sources"
    assert "flood_10y" not in child["modules"] and child["version_changes"] == []
    r = client.post(f"/api/runs/{first}/rerun", json={"kind": "latest"})
    assert r.status_code == 201 and r.json()["rerun_kind"] == "latest"
    assert client.get(f"/api/runs/{r.json()['id']}/changes").json() == []  # nothing new
    journal = {x["id"]: x for x in client.get("/api/runs").json()["runs"]}
    assert journal[child["id"]]["parent_run_id"] == first
    bad = client.post(f"/api/runs/{first}/rerun", json={"kind": "sources", "sources": ["nope"]})
    assert bad.status_code == 422
    assert client.post("/api/runs/999999/rerun", json={"kind": "latest"}).status_code == 404


def test_api_groups(client, first):
    g = client.get(f"/api/runs/{first}/groups").json()
    assert g["id"] is None and [x["name"] for x in g["groups"]] == ["P-1", "P-2", "Bez projekta"]
    body = {"groups": [{"id": "x", "name": "Kopā"}], "assignment": {"0": "x", "1": "x"}}
    r = client.post(f"/api/runs/{first}/groups", json=body)
    assert r.status_code == 201 and r.json()["id"] is not None
    assert client.get(f"/api/runs/{first}/groups").json()["assignment"] == {"0": "x", "1": "x"}
    # Žurnāls follows the grouping, also in its project search
    journal = {x["id"]: x for x in client.get("/api/runs", params={"limit": 200}).json()["runs"]}
    assert journal[first]["projects"] == ["Kopā"]
    found = client.get("/api/runs", params={"project": "kop", "limit": 200}).json()["runs"]
    assert first in [x["id"] for x in found]
    gone = client.get("/api/runs", params={"project": "P-2", "limit": 200}).json()["runs"]
    assert first not in [x["id"] for x in gone]
    bad = client.post(f"/api/runs/{first}/groups", json={"groups": [], "assignment": {"0": "none"}})
    assert bad.status_code == 422
    report = client.get(f"/api/runs/{first}/report.xlsx")
    ws = load_workbook(io.BytesIO(report.content))["Kopsavilkums"]
    assert ws.cell(2, 3).value == "Kopā"
    old = client.get(f"/api/runs/{first}/report.xlsx", params={"grouping": 999999})
    assert old.status_code == 404


# --- Žurnāls filters and test origins ----------------------------------------------------------


def test_journal_hides_test_runs_and_pages(client, conn, monkeypatch):
    monkeypatch.setenv("KADASTRS_RUN_ORIGIN", "e2e")  # what api-e2e / api-screens set
    hidden = client.post("/api/batch", data={"text": "09000050400"}).json()["id"]
    monkeypatch.delenv("KADASTRS_RUN_ORIGIN")
    shown = client.post("/api/batch", data={"text": "09000050400"}).json()["id"]
    assert (
        conn.execute("SELECT origin FROM work.run WHERE id = %s", (hidden,)).fetchone()[0] == "e2e"
    )

    assert shown in journal_ids(client) and hidden not in journal_ids(client)
    with_tests = client.get("/api/runs", params={"tests": True, "limit": 200}).json()
    assert hidden in [x["id"] for x in with_tests["runs"]]
    assert next(x for x in with_tests["runs"] if x["id"] == hidden)["test"] is True
    # paging: 1 per page, newest first, no overlap
    total = client.get("/api/runs", params={"tests": True, "limit": 1}).json()["total"]
    first_page = journal_ids(client, tests=True, limit=1)
    second_page = journal_ids(client, tests=True, limit=1, offset=1)
    assert total == with_tests["total"] and first_page[0] > second_page[0]


def test_journal_date_range(client, conn):
    rid = client.post("/api/batch", data={"text": "09000050400"}).json()["id"]
    day = conn.execute(
        "SELECT (created_at AT TIME ZONE 'Europe/Riga')::date FROM work.run WHERE id = %s", (rid,)
    ).fetchone()[0]

    assert rid in journal_ids(client, date_from=day.isoformat(), date_to=day.isoformat())
    assert rid not in journal_ids(client, date_to=(day - timedelta(days=1)).isoformat())
    assert rid not in journal_ids(client, date_from=(day + timedelta(days=1)).isoformat())
    assert client.get("/api/runs", params={"date_from": "nav datums"}).status_code == 422


def test_project_search_takes_like_characters_literally(client):
    rid = client.post("/api/batch", data={"text": "vertiba;projekta_nr\n09000050400;A_1%"}).json()[
        "id"
    ]
    hits = client.get("/api/runs", params={"project": "A_1%", "limit": 200}).json()["runs"]
    assert rid in [x["id"] for x in hits]
    miss = client.get("/api/runs", params={"project": "AX1", "limit": 200}).json()["runs"]
    assert rid not in [x["id"] for x in miss]


def test_preview_counts_property_units(client):
    units = client.get("/api/property/09000360220/units").json()["units"]
    rows = client.post("/api/batch/preview", data={"text": "09000360220\n09000050400"}).json()[
        "rows"
    ]
    by = {t["type"]: t["units"] for r in rows for t in r["targets"]}
    assert by["property"] == len(units) and by["parcel"] == 1
