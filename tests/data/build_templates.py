"""Build the two Phase 0 templates for colleagues (PLAN.md §2.4 steps 4 and 5).

    .venv/Scripts/python -I tests/data/build_templates.py

Writes tests/data/test_set_template.xlsx and tests/data/laika_uzskaite.xlsx.
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1]))  # the repo root, for engine.batch (run with -I)

from engine.batch.headers import INPUT_COLUMNS, TEST_SET_COLUMNS  # noqa: E402
from engine.batch.template import LEGEND, style_headers  # noqa: E402
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="305496")
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, columns, widths):
    for i, (name, width) in enumerate(zip(columns, widths), 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def instructions(ws, lines):
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = WRAP
        if i == 1 or line.endswith(":"):
            c.font = Font(bold=True, size=13 if i == 1 else 11)


def test_set():
    wb = Workbook()
    ws = wb.active
    ws.title = "Testa_kopa"
    # readable Latvian headers (old names still accepted, engine/batch/headers.py)
    by_name = {c.name: c for c in INPUT_COLUMNS + TEST_SET_COLUMNS}
    names = ["projekta_nr", "vertiba", "adrese", "veids", "pieteicejs_reg_nr",
             "m1_kadastra_dati", "m2_adreses_atbilstiba", "m3_aizsargajama_daba", "m4_pludi_10_gadi",
             "m5_buvniecibas_lietas", "m6_apgrutinajumi", "parbaudes_datums", "avots_piezimes"]
    widths = [16, 30, 32, 26, 24, 34, 28, 34, 30, 34, 34, 16, 44]
    style_headers(ws, [(by_name[n], w) for n, w in zip(names, widths)], rows=500)
    hint = "‹Ir: … / Nav / Nav datu / Nav pārklājuma›"
    examples = [
        ["PIEMĒRS-1 (izdzēst)", "05000041006", "Tautas iela 11, Daugavpils", "īpašums", "90000077325", hint, hint, hint, hint, hint, hint, None,
         "Piemērs no sākotnējiem datiem: Daugavpils, Tautas iela 11. Atbildes aizpilda kolēģis."],
        ["PIEMĒRS-2 (izdzēst)", "09000360220", None, None, None, hint, hint, hint, hint, hint, hint, None,
         "Piemērs no sākotnējiem datiem: Jelgava. Atbildes aizpilda kolēģis."],
    ]
    for r, row in enumerate(examples, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill, cell.alignment = EXAMPLE_FILL, WRAP
    date_col = get_column_letter(names.index("parbaudes_datums") + 1)
    for r in range(2, 502):
        ws[f"{date_col}{r}"].number_format = "yyyy-mm-dd"
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                        error="Ievadiet datumu (GGGG-MM-DD).", errorTitle="Datums")
    ws.add_data_validation(dv)
    dv.add(f"{date_col}2:{date_col}501")

    info = wb.create_sheet("Instrukcija", 0)
    instructions(info, [
        "Testa kopa: manuālās pārbaudes atbildes",
        "Mērķis: salīdzināt rīka rezultātus ar jūsu manuālās pārbaudes atbildēm. Vajag 14–20 reālus projektus, "
        "vēlams Jelgavā, Jūrmalā vai Daugavpilī (pilota teritorija).",
        "",
        "Kā aizpildīt:",
        "• Viena rinda = viens objekts. Ja projektā ir vairāki objekti, katram sava rinda ar to pašu projekta_nr.",
        f"• {by_name['projekta_nr'].header}: projekta numurs (brīvs teksts).",
        f"• {by_name['vertiba'].header} (obligāta): kadastra numurs (11 cipari), kadastra apzīmējums (11, 14 vai 17 cipari) vai adrese. "
        "Kolonnas A–E ir teksta formātā: ievadiet ciparus ar sākuma nullēm (piem., 05000041006).",
        f"• {by_name['adrese'].header}: adrese, kā tā norādīta pieteikumā. Rīks pārbauda, vai tā atbilst "
        "objektam (M2); bez tās M2 netiek salīdzināts.",
        f"• {by_name['veids'].header}: izvēlieties no saraksta, ja 11 ciparu numurs var būt gan "
        "īpašuma numurs, gan citam īpašumam piederošas zemes vienības apzīmējums; bez tā rīks pārbauda abus.",
        f"• {by_name['pieteicejs_reg_nr'].header}: pieteicēja juridiskās personas reģistrācijas numurs (ja ir). "
        "Fizisko personu vārdus un personas kodus neievadiet.",
        "• M1–M6: jūsu manuālās pārbaudes atbilde katram modulim. Sāciet ar vienu no: "
        "«Ir: …» (atrasts, īsi ko), «Nav», «Nav datu» (objekts nav atrasts / nav ģeometrijas), "
        "«Nav pārklājuma» (avotā šai teritorijai nav datu, piem., nav plūdu kartes). Pēc tam īss paskaidrojums.",
        f"• {by_name['parbaudes_datums'].header}: datums, kad veicāt pārbaudi (GGGG-MM-DD). Svarīgi, jo dati laika gaitā mainās.",
        f"• {by_name['avots_piezimes'].header}: kurus portālus izmantojāt un citas piezīmes.",
        "",
        *LEGEND,
        "",
        "Moduļi:",
        "m1 Kadastra dati: zemes vienības / būves, platība, lietošanas mērķis, būves uz zemes vienības.",
        "m2 Adreses atbilstība: vai adrese pieteikumā (kolonna „Adrese pieteikumā”) atbilst kadastra objektam («Neatbilst»).",
        "m3 Aizsargājamā daba: ĪADT, mikroliegums, mikrolieguma buferzona, dabas piemineklis, aizsargājams biotops "
        "(katru atsevišķi).",
        "m4 Plūdi: 3. cikla pavasara palu karte, reizi 10 gados (10 % varbūtība). Vai objekts ir plūdu zonā, "
        "aptuvenais pārklājums.",
        "m5 Būvniecības lietas (BIS): lietas uz šī objekta; norādiet BIS lietas numurus un vai tās ir tieši uz "
        "objekta vai blakus.",
        "m6 Apgrūtinājumi: kadastrā reģistrētie apgrūtinājumi (kods vai nosaukums).",
        "",
        "Lūdzam testa kopā iekļaut:",
        "• vismaz vienu zemes vienību, kas atrodas 10 gadu plūdu zonā;",
        "• vismaz vienu objektu ārpus plūdu karšu pārklājuma (teritorijā, kurai plūdu karte nav izstrādāta);",
        "• ja iespējams: objektu ar būvniecības lietu, objektu dabas teritorijā un objektu ar vairākām "
        "zemes vienībām vai būvēm.",
        "",
        "Dzeltenās rindas ir piemēri: izdzēsiet tās vai aizstājiet ar saviem datiem.",
    ])
    wb.save(HERE / "test_set_template.xlsx")


def timing():
    wb = Workbook()
    ws = wb.active
    ws.title = "Laika_uzskaite"
    cols = ["objekts", "datums", "sakums", "beigas", "minutes", "izmantotie_portali", "piezimes"]
    header(ws, cols, [30, 14, 10, 10, 10, 44, 44])
    for r in range(2, 202):
        ws[f"A{r}"].number_format = "@"
        ws[f"B{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"].number_format = "hh:mm"
        ws[f"D{r}"].number_format = "hh:mm"
        ws[f"E{r}"] = f'=IF(AND(C{r}<>"",D{r}<>""),ROUND((D{r}-C{r})*1440,0),"")'
        ws[f"E{r}"].number_format = "0"
    ws["H1"], ws["H2"] = "Vidēji minūtes uz objektu", '=IFERROR(AVERAGE(E2:E201),"")'
    ws["H1"].font = Font(bold=True)
    ws.column_dimensions["H"].width = 26
    info = wb.create_sheet("Instrukcija", 0)
    instructions(info, [
        "Laika uzskaite: manuālās pārbaudes bāzes līnija",
        "Mērķis: izmērīt, cik ilgi šobrīd aizņem viena objekta pārbaude. Ar šo skaitli salīdzināsim rīku.",
        "",
        "Kā aizpildīt:",
        "• Pārbaudiet 5 objektus tāpat kā ikdienā, un katram aizpildiet vienu rindu.",
        "• objekts: kadastra numurs, apzīmējums vai adrese (teksta formātā, ar sākuma nullēm).",
        "• datums: pārbaudes datums.",
        "• sakums: laiks (SS:MM), kad atverat pirmo portālu šim objektam.",
        "• beigas: laiks (SS:MM), kad visas atbildes ir pierakstītas.",
        "• minutes: aprēķinās automātiski.",
        "• izmantotie_portali: visi izmantotie portāli (piem., Kadastrs.lv, Ozols, geolatvija.lv, BIS, LVĢMC plūdu karte).",
        "• piezimes: kas aizņēma visvairāk laika, problēmas, pārtraukumi.",
        "",
        "Pārbaudē iekļaujiet tos pašus moduļus kā testa kopā: kadastra dati, adreses atbilstība, aizsargājamā daba, "
        "plūdi (10 gadu), būvniecības lietas, apgrūtinājumi.",
        "Pārtraukumus neieskaitiet: ja pārtraucāt darbu, pierakstiet to piezīmēs.",
    ])
    wb.save(HERE / "laika_uzskaite.xlsx")


if __name__ == "__main__":
    test_set()
    timing()
    print("written:", HERE / "test_set_template.xlsx", HERE / "laika_uzskaite.xlsx")
