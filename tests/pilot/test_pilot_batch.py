"""Phase 5 pilot batch: 1,000 real input rows (parcels, buildings, addresses) ->
run -> Excel report, timed against the PLAN.md §9 target (1,000 rows < 5 min).
The run is stored in the dev DB like any other run (kept, D3)."""

import io
import time

import psycopg
import pytest
from openpyxl import load_workbook

from conftest import TEST_DATABASE_URL
from engine.batch.plan import run_batch
from engine.batch.reader import read_text
from engine.checks import runner
from reports import excel

pytestmark = [pytest.mark.pilot, pytest.mark.stack]

BATCH_1000_LIMIT_S = 300  # PLAN.md §9: a 1,000-row batch finishes in < 5 min


@pytest.fixture
def conn(live):
    with psycopg.connect(TEST_DATABASE_URL, autocommit=True) as c:
        yield c


def test_thousand_real_rows_within_the_time_limit(conn, live):
    t = live["cadastre_text"]
    # a stable spread over the three cities: parcels, buildings and addresses
    parcels = [
        r[0]
        for r in conn.execute(
            "SELECT code FROM ref.parcel WHERE dataset_version_id = %s"
            " ORDER BY md5(code) LIMIT 600",
            (t,),
        )
    ]
    buildings = [
        r[0]
        for r in conn.execute(
            "SELECT code FROM ref.building WHERE dataset_version_id = %s AND NOT not_exist"
            " ORDER BY md5(code) LIMIT 300",
            (t,),
        )
    ]
    addresses = [
        r[0]
        for r in conn.execute(
            "SELECT full_text FROM ref.address WHERE dataset_version_id = %s AND status = 'EKS'"
            " ORDER BY md5(code) LIMIT 100",
            (live["varis_addresses"],),
        )
    ]
    values = parcels + buildings + addresses
    assert len(values) == 1000
    text = "vertiba;projekta_nr\n" + "\n".join(
        f'"{v}";P-{i // 10:03d}' for i, v in enumerate(values)
    )

    t0 = time.perf_counter()
    rows = read_text(text, 1000)
    run_id = run_batch(conn, rows, origin="test")  # hidden from Žurnāls by default
    report = excel.to_bytes(excel.load(conn, run_id))
    elapsed = time.perf_counter() - t0

    inputs = runner.stored_inputs(conn, run_id)
    checked = sum(1 for i in inputs if i["target_nos"])
    n_targets = conn.execute(
        "SELECT count(*) FROM work.run_target WHERE run_id = %s", (run_id,)
    ).fetchone()[0]
    wb = load_workbook(io.BytesIO(report), read_only=True)
    print(
        f"\npilot batch run {run_id}: 1,000 rows, {checked} checked, {n_targets} objects,"
        f" {elapsed:.1f} s ({elapsed * 1000 / len(values):.0f} ms per row),"
        f" report {len(report)} bytes"
    )
    assert len(inputs) == 1000
    assert checked >= 950  # real designations and current addresses resolve
    assert wb["Kopsavilkums"].max_row >= n_targets + 1
    assert elapsed < BATCH_1000_LIMIT_S
